# ============================================================
# CVSmart V11 - VERSION FINALE COMPLÈTE
# Toutes les fonctionnalités intégrées
# ============================================================

import io
import os
import re
import json
import shutil
import hashlib
import unicodedata
import warnings
from pathlib import Path
from datetime import datetime
from typing import Dict, List, Tuple, Optional

warnings.filterwarnings('ignore')

import fitz
import docx
import pandas as pd
import numpy as np
import streamlit as st
import plotly.express as px
import plotly.graph_objects as go

from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.naive_bayes import MultinomialNB
from sklearn.pipeline import Pipeline
from gestionnaire_email import GestionnaireEmail, tester_connexion
import ui
from html import escape

# ============================================================
# CONFIGURATION
# ============================================================
st.set_page_config(
    page_title="CVSmart - L'intelligence qui trie vos talents",
    page_icon=ui.logo_icone(),
    layout="wide",
    initial_sidebar_state="expanded"
)
VERSION_APP = "1.2"

# ============================================================
# 50 DOMAINES AVEC MÉTIERS
# ============================================================
CATEGORIES = {
    'Informatique': {
        'emoji': '💻', 'couleur': '#2B3F8C',
        'metiers': {
            'Développeur': ['développeur', 'programmeur', 'full stack', 'front-end', 'back-end', 'python', 'java', 'javascript', 'php', 'react', 'node'],
            'Administrateur réseau': ['administrateur réseau', 'admin réseau', 'réseau', 'ccna', 'cisco', 'windows server'],
            'DevOps': ['devops', 'docker', 'kubernetes', 'ci/cd', 'jenkins', 'ansible'],
            'Data Scientist': ['data scientist', 'data analyst', 'machine learning', 'big data'],
            'Cybersécurité': ['cybersécurité', 'sécurité informatique', 'pentest', 'ceh'],
            'Chef de projet IT': ['chef de projet informatique', 'scrum master', 'product owner'],
            'Support IT': ['support informatique', 'helpdesk', 'technicien informatique'],
        },
        'mots_cles': ['informatique', 'logiciel', 'application', 'web', 'mobile', 'cloud', 'sql', 'api', 'git', 'linux']
    },
    'Génie Civil': {
        'emoji': '🏗️', 'couleur': '#FF6B00',
        'metiers': {
            'Ingénieur civil': ['ingénieur civil', 'ingénieur génie civil', 'ingénieur structure'],
            'Conducteur de travaux': ['conducteur de travaux', 'chef de chantier'],
            'Topographe': ['topographe', 'géomètre'],
            'Dessinateur projeteur': ['dessinateur', 'projeteur', 'autocad'],
        },
        'mots_cles': ['génie civil', 'chantier', 'béton', 'structure', 'construction', 'plans']
    },
    'Génie Électrique': {
        'emoji': '⚡', 'couleur': '#FFD700',
        'metiers': {
            'Ingénieur électrique': ['ingénieur électrique', 'ingénieur électricité'],
            'Électricien': ['électricien', 'électricité', 'câblage', 'installation électrique'],
            'Technicien électricité': ['technicien électricité', 'technicien électrique'],
            'Automaticien': ['automaticien', 'automate', 'plc', 'scada'],
        },
        'mots_cles': ['électrique', 'électricité', 'courant', 'tension', 'circuit']
    },
    'Génie Mécanique': {
        'emoji': '⚙️', 'couleur': '#8B00FF',
        'metiers': {
            'Ingénieur mécanique': ['ingénieur mécanique'],
            'Technicien mécanique': ['technicien mécanique', 'mécanicien industriel'],
            'Dessinateur industriel': ['dessinateur industriel', 'solidworks', 'catia'],
            'Responsable maintenance': ['responsable maintenance', 'maintenance industrielle'],
        },
        'mots_cles': ['mécanique', 'machine', 'moteur', 'usinage', 'fabrication']
    },
    'Génie des Procédés': {
        'emoji': '🧪', 'couleur': '#E63946',
        'metiers': {
            'Ingénieur procédés': ['ingénieur procédés', 'génie des procédés', 'procédés industriels'],
            'Ingénieur chimiste': ['ingénieur chimiste', 'chimie industrielle'],
            'Responsable production': ['responsable production', 'chef de production'],
            'Technicien procédés': ['technicien procédés', 'opérateur procédés'],
        },
        'mots_cles': ['procédés', 'chimie', 'pétrole', 'raffinage', 'industriel']
    },
    'Mécatronique': {
        'emoji': '🤖', 'couleur': '#00C4CC',
        'metiers': {
            'Ingénieur mécatronique': ['ingénieur mécatronique', 'mécatronique'],
            'Roboticien': ['roboticien', 'robotique', 'robot'],
            'Technicien automatisme': ['automatisme', 'automate', 'plc'],
        },
        'mots_cles': ['mécatronique', 'robot', 'automatisation', 'électronique']
    },
    'Électrotechnique': {
        'emoji': '🔌', 'couleur': '#FFAA00',
        'metiers': {
            'Ingénieur électrotechnique': ['électrotechnique', 'ingénieur électrotechnique'],
            'Technicien électrotechnique': ['technicien électrotechnique'],
        },
        'mots_cles': ['électrotechnique', 'moteur électrique', 'transformateur']
    },
    'Télécommunications': {
        'emoji': '📡', 'couleur': '#7B68EE',
        'metiers': {
            'Ingénieur télécom': ['ingénieur télécom', 'ingénieur télécommunications'],
            'Technicien réseau': ['technicien réseau', 'technicien télécom'],
            'Installateur': ['installateur antenne', 'installation réseau'],
        },
        'mots_cles': ['télécom', 'télécommunications', 'fibre', 'antenne', '4g', '5g']
    },
    'BTP et Construction': {
        'emoji': '🔨', 'couleur': '#A0522D',
        'metiers': {
            'Maçon': ['maçon', 'maçonnerie'],
            'Coffreur': ['coffreur', 'coffrage'],
            'Ferrailleur': ['ferrailleur', 'ferraillage'],
            'Chef de chantier': ['chef de chantier'],
            'Peintre bâtiment': ['peintre bâtiment', 'peinture'],
            'Carreleur': ['carreleur', 'carrelage'],
        },
        'mots_cles': ['btp', 'bâtiment', 'chantier', 'maçonnerie']
    },
    'Énergie et Pétrole': {
        'emoji': '🛢️', 'couleur': '#2F4F4F',
        'metiers': {
            'Ingénieur pétrole': ['ingénieur pétrole', 'ingénieur pétrolier'],
            'Ingénieur énergie': ['ingénieur énergie', 'énergies renouvelables'],
            'Technicien pétrole': ['technicien pétrole', 'technicien forage'],
        },
        'mots_cles': ['pétrole', 'gaz', 'énergie', 'forage', 'raffinage', 'solaire']
    },
    'Comptabilité': {
        'emoji': '💰', 'couleur': '#E63946',
        'metiers': {
            'Comptable': ['comptable', 'comptabilité', 'tenue de comptes'],
            'Auditeur': ['auditeur', 'audit'],
            'Contrôleur de gestion': ['contrôleur de gestion', 'contrôle de gestion'],
            'Expert-comptable': ['expert-comptable', 'expert comptable'],
            'Assistant comptable': ['assistant comptable', 'aide-comptable'],
        },
        'mots_cles': ['comptabilité', 'bilan', 'grand livre', 'sage', 'écritures']
    },
    'Finance': {
        'emoji': '💵', 'couleur': '#00C4CC',
        'metiers': {
            'Analyste financier': ['analyste financier', 'analyse financière'],
            'Directeur financier': ['directeur financier', 'daf'],
            'Contrôleur financier': ['contrôleur financier'],
            'Trésorier': ['trésorier', 'trésorerie'],
        },
        'mots_cles': ['finance', 'budget', 'investissement', 'portefeuille', 'risque']
    },
    'Banque et Assurance': {
        'emoji': '🏦', 'couleur': '#2B3F8C',
        'metiers': {
            'Chargé de clientèle': ['chargé de clientèle', 'conseiller clientèle'],
            'Banquier': ['banquier', 'banque'],
            'Assureur': ['assureur', 'assurance'],
            'Analyste crédit': ['analyste crédit', 'analyse crédit'],
        },
        'mots_cles': ['banque', 'assurance', 'crédit', 'placement', 'microfinance']
    },
    'Ressources Humaines': {
        'emoji': '👥', 'couleur': '#8B00FF',
        'metiers': {
            'Chargé RH': ['chargé rh', 'chargé ressources humaines'],
            'Recruteur': ['recruteur', 'recrutement', 'talent acquisition'],
            'Gestionnaire de paie': ['gestionnaire paie', 'paie', 'payroll'],
            'Responsable formation': ['responsable formation'],
            'DRH': ['drh', 'directeur ressources humaines'],
        },
        'mots_cles': ['rh', 'ressources humaines', 'recrutement', 'paie', 'sirh']
    },
    'Marketing': {
        'emoji': '📢', 'couleur': '#FF6B00',
        'metiers': {
            'Marketing digital': ['marketing digital', 'digital marketing', 'community manager'],
            'Chef de produit': ['chef de produit', 'product manager'],
            'Brand manager': ['brand manager', 'branding', 'marque'],
            'SEO/SEA': ['seo', 'sem', 'référencement'],
            'Chargé de communication': ['chargé de communication', 'communication'],
        },
        'mots_cles': ['marketing', 'communication', 'digital', 'campagne', 'publicité']
    },
    'Commercial et Vente': {
        'emoji': '🤝', 'couleur': '#FF0055',
        'metiers': {
            'Commercial': ['commercial', 'commerciale', 'représentant'],
            'Business Developer': ['business developer', 'business development'],
            'Responsable des ventes': ['responsable des ventes', 'chef des ventes'],
            'Téléconseiller': ['téléconseiller', 'télémarketing'],
        },
        'mots_cles': ['vente', 'commerce', 'négociation', 'client', 'prospection']
    },
    'Logistique et Transport': {
        'emoji': '🚚', 'couleur': '#FFAA00',
        'metiers': {
            'Logisticien': ['logisticien', 'logistique'],
            'Chauffeur': ['chauffeur', 'conducteur', 'permis'],
            'Magasinier': ['magasinier', 'magasin', 'gestion de stock'],
            'Gestionnaire de stock': ['gestionnaire stock', 'gestion des stocks'],
            'Transitaire': ['transitaire', 'transit', 'douane'],
            'Chef de dépôt': ['chef de dépôt', 'responsable dépôt'],
        },
        'mots_cles': ['logistique', 'transport', 'supply chain', 'approvisionnement', 'distribution']
    },
    'Achats et Approvisionnement': {
        'emoji': '📦', 'couleur': '#E63946',
        'metiers': {
            'Acheteur': ['acheteur', 'acheteuse', 'achats'],
            'Responsable achats': ['responsable achats'],
            'Approvisionneur': ['approvisionneur', 'approvisionnement'],
        },
        'mots_cles': ['achat', 'approvisionnement', 'fournisseur', 'commande']
    },
    'Administration': {
        'emoji': '📋', 'couleur': '#666666',
        'metiers': {
            'Secrétaire': ['secrétaire', 'secrétariat'],
            'Assistant de direction': ['assistant de direction', 'assistante de direction'],
            'Office Manager': ['office manager'],
            'Agent administratif': ['agent administratif', 'administration'],
        },
        'mots_cles': ['administration', 'secrétariat', 'assistant', 'planning', 'agenda']
    },
    'Direction Générale': {
        'emoji': '👔', 'couleur': '#2B3F8C',
        'metiers': {
            'Directeur Général': ['directeur général', 'dg'],
            'Directeur Général Adjoint': ['directeur général adjoint', 'dga'],
            'Président': ['président', 'pdg', 'ceo'],
            'Directeur': ['directeur', 'directrice'],
            'Manager': ['manager', 'responsable'],
        },
        'mots_cles': ['direction', 'général', 'président', 'manager', 'responsable']
    },
    'Santé et Médical': {
        'emoji': '⚕️', 'couleur': '#FF0055',
        'metiers': {
            'Médecin': ['médecin', 'docteur', 'doctorat en médecine'],
            'Infirmier': ['infirmier', 'infirmière'],
            'Sage-femme': ['sage-femme', 'sage femme'],
            'Kinésithérapeute': ['kinésithérapeute', 'kiné'],
        },
        'mots_cles': ['santé', 'médical', 'hôpital', 'patient', 'soin']
    },
    'Pharmacie': {
        'emoji': '💊', 'couleur': '#E63946',
        'metiers': {
            'Pharmacien': ['pharmacien', 'pharmacie'],
            'Préparateur en pharmacie': ['préparateur pharmacie'],
        },
        'mots_cles': ['pharmacie', 'médicament', 'pharmaceutique']
    },
    'Paramédical': {
        'emoji': '🩺', 'couleur': '#00C4CC',
        'metiers': {
            'Laborantin': ['laborantin', 'laboratoire'],
            'Radiologue': ['radiologue', 'imagerie'],
            'Technicien laboratoire': ['technicien laboratoire'],
        },
        'mots_cles': ['paramédical', 'laboratoire', 'analyse', 'imagerie']
    },
    'Action Sociale': {
        'emoji': '🤲', 'couleur': '#FF6B00',
        'metiers': {
            'Assistant social': ['assistant social', 'assistante sociale'],
            'Éducateur spécialisé': ['éducateur', 'éducateur spécialisé'],
        },
        'mots_cles': ['social', 'aide', 'humanitaire', 'ong']
    },
    'Vétérinaire': {
        'emoji': '🐕', 'couleur': '#E63946',
        'metiers': {
            'Vétérinaire': ['vétérinaire', 'veterinaire'],
            'Assistant vétérinaire': ['assistant vétérinaire'],
        },
        'mots_cles': ['vétérinaire', 'animaux', 'élevage']
    },
    'Enseignement': {
        'emoji': '🎓', 'couleur': '#2B3F8C',
        'metiers': {
            'Enseignant': ['enseignant', 'professeur', 'instituteur'],
            'Professeur': ['professeur', 'maître de conférences'],
            'Directeur d\'école': ['directeur école', 'proviseur'],
        },
        'mots_cles': ['enseignement', 'école', 'université', 'cours', 'élève']
    },
    'Formation professionnelle': {
        'emoji': '📚', 'couleur': '#8B00FF',
        'metiers': {
            'Formateur': ['formateur', 'formatrice'],
            'Coach': ['coach', 'coaching'],
            'Responsable formation': ['responsable formation'],
        },
        'mots_cles': ['formation', 'formateur', 'coaching', 'pédagogie']
    },
    'Recherche scientifique': {
        'emoji': '🔬', 'couleur': '#00C4CC',
        'metiers': {
            'Chercheur': ['chercheur', 'chercheuse'],
            'Assistant de recherche': ['assistant recherche'],
        },
        'mots_cles': ['recherche', 'scientifique', 'laboratoire', 'publication']
    },
    'Droit et Juridique': {
        'emoji': '⚖️', 'couleur': '#2B3F8C',
        'metiers': {
            'Juriste': ['juriste', 'juridique'],
            'Avocat': ['avocat', 'avocate'],
            'Assistant juridique': ['assistant juridique'],
        },
        'mots_cles': ['droit', 'juridique', 'loi', 'contrat', 'contentieux']
    },
    'Notariat': {
        'emoji': '📜', 'couleur': '#666666',
        'metiers': {
            'Notaire': ['notaire', 'notariat'],
            'Clerc de notaire': ['clerc de notaire'],
        },
        'mots_cles': ['notaire', 'notariat', 'acte']
    },
    'Communication et Journalisme': {
        'emoji': '📰', 'couleur': '#FF6B00',
        'metiers': {
            'Journaliste': ['journaliste', 'journalisme'],
            'Rédacteur': ['rédacteur', 'rédactrice', 'rédaction'],
            'Attaché de presse': ['attaché de presse'],
        },
        'mots_cles': ['journalisme', 'presse', 'média', 'rédaction', 'reportage']
    },
    'Art et Design': {
        'emoji': '🎨', 'couleur': '#FF0055',
        'metiers': {
            'Graphiste': ['graphiste', 'graphisme', 'photoshop', 'illustrator'],
            'Designer': ['designer', 'webdesigner', 'ux', 'ui'],
            'Directeur artistique': ['directeur artistique', 'da'],
        },
        'mots_cles': ['design', 'graphique', 'créatif', 'art', 'illustration']
    },
    'Architecture et Urbanisme': {
        'emoji': '🏛️', 'couleur': '#8B00FF',
        'metiers': {
            'Architecte': ['architecte', 'architecture'],
            'Urbaniste': ['urbaniste', 'urbanisme'],
            'Paysagiste': ['paysagiste'],
        },
        'mots_cles': ['architecture', 'urbanisme', 'plans', 'maquette']
    },
    'Audiovisuel et Cinéma': {
        'emoji': '🎬', 'couleur': '#FFAA00',
        'metiers': {
            'Réalisateur': ['réalisateur', 'réalisatrice'],
            'Monteur': ['monteur', 'montage'],
            'Cadreur': ['cadreur', 'caméraman'],
            'Photographe': ['photographe', 'photographie'],
        },
        'mots_cles': ['audiovisuel', 'cinéma', 'vidéo', 'film', 'tournage']
    },
    'Mode et Textile': {
        'emoji': '👗', 'couleur': '#FF0055',
        'metiers': {
            'Styliste': ['styliste', 'stylisme'],
            'Couturier': ['couturier', 'couturière', 'couture'],
            'Modéliste': ['modéliste'],
        },
        'mots_cles': ['mode', 'textile', 'couture', 'vêtement', 'stylisme']
    },
    'Hôtellerie et Restauration': {
        'emoji': '🍽️', 'couleur': '#FF6B00',
        'metiers': {
            'Cuisinier': ['cuisinier', 'cuisinière', 'chef cuisinier'],
            'Chef de cuisine': ['chef de cuisine'],
            'Serveur': ['serveur', 'serveuse'],
            'Réceptionniste': ['réceptionniste', 'réception'],
            'Barman': ['barman', 'barmaid'],
        },
        'mots_cles': ['hôtel', 'restaurant', 'cuisine', 'service', 'tourisme']
    },
    'Tourisme': {
        'emoji': '✈️', 'couleur': '#00C4CC',
        'metiers': {
            'Guide touristique': ['guide touristique', 'guide'],
            'Agent de voyage': ['agent de voyage', 'agence de voyage'],
        },
        'mots_cles': ['tourisme', 'voyage', 'guide', 'agence']
    },
    'Agriculture': {
        'emoji': '🌾', 'couleur': '#E63946',
        'metiers': {
            'Agronome': ['agronome', 'agronomie'],
            'Agriculteur': ['agriculteur', 'agricultrice', 'agriculture'],
            'Technicien agricole': ['technicien agricole'],
        },
        'mots_cles': ['agriculture', 'culture', 'champ', 'récolte', 'élevage']
    },
    'Agroalimentaire': {
        'emoji': '🍞', 'couleur': '#FFAA00',
        'metiers': {
            'Ingénieur agroalimentaire': ['agroalimentaire', 'ingénieur agro'],
            'Technicien qualité': ['technicien qualité', 'contrôle qualité'],
        },
        'mots_cles': ['agroalimentaire', 'alimentaire', 'transformation', 'qualité']
    },
    'Environnement': {
        'emoji': '🌍', 'couleur': '#E63946',
        'metiers': {
            'Ingénieur environnement': ['ingénieur environnement'],
            'Écologiste': ['écologiste', 'écologie'],
        },
        'mots_cles': ['environnement', 'écologie', 'durable', 'déchets']
    },
    'QHSE': {
        'emoji': '🦺', 'couleur': '#FF6B00',
        'metiers': {
            'Responsable QHSE': ['qhse', 'responsable qhse', 'hse'],
            'Animateur sécurité': ['animateur sécurité', 'sécurité'],
            'Contrôleur qualité': ['contrôleur qualité', 'qualité'],
        },
        'mots_cles': ['qhse', 'hse', 'sécurité', 'hygiène', 'qualité', 'environnement']
    },
    'Artisanat': {
        'emoji': '🔧', 'couleur': '#A0522D',
        'metiers': {
            'Artisan': ['artisan', 'artisanat'],
            'Soudeur': ['soudeur', 'soudure'],
            'Bijoutier': ['bijoutier', 'bijouterie'],
            'Cordonnier': ['cordonnier', 'cordonnerie'],
        },
        'mots_cles': ['artisanat', 'artisan', 'manuel', 'fabrication']
    },
    'Menuiserie et Charpente': {
        'emoji': '🪚', 'couleur': '#A0522D',
        'metiers': {
            'Menuisier': ['menuisier', 'menuiserie', 'bois'],
            'Charpentier': ['charpentier', 'charpente'],
            'Ébéniste': ['ébéniste', 'ébénisterie'],
        },
        'mots_cles': ['menuiserie', 'bois', 'charpente', 'meuble']
    },
    'Plomberie': {
        'emoji': '🔧', 'couleur': '#2B3F8C',
        'metiers': {
            'Plombier': ['plombier', 'plomberie'],
            'Installateur sanitaire': ['installation sanitaire'],
        },
        'mots_cles': ['plomberie', 'tuyau', 'sanitaire', 'eau']
    },
    'Mécanique Auto': {
        'emoji': '🚗', 'couleur': '#FF0055',
        'metiers': {
            'Mécanicien auto': ['mécanicien auto', 'mécanique auto'],
            'Carrossier': ['carrossier', 'carrosserie'],
            'Électricien auto': ['électricien auto'],
            'Vidangeur': ['vidangeur', 'vidange', 'graissage'],
        },
        'mots_cles': ['automobile', 'voiture', 'mécanique', 'moteur', 'garage']
    },
    'Métallurgie et Soudure': {
        'emoji': '⚒️', 'couleur': '#666666',
        'metiers': {
            'Soudeur': ['soudeur', 'soudure', 'soudage'],
            'Chaudronnier': ['chaudronnier', 'chaudronnerie'],
            'Tôlier': ['tôlier', 'tôlerie'],
        },
        'mots_cles': ['métallurgie', 'métal', 'soudure', 'acier']
    },
    'Sécurité': {
        'emoji': '🛡️', 'couleur': '#2B3F8C',
        'metiers': {
            'Agent de sécurité': ['agent de sécurité', 'vigile'],
            'Gardien': ['gardien', 'gardiennage'],
            'Chef de sécurité': ['chef de sécurité', 'responsable sécurité'],
        },
        'mots_cles': ['sécurité', 'gardiennage', 'surveillance', 'protection']
    },
    'Nettoyage et Entretien': {
        'emoji': '🧹', 'couleur': '#00C4CC',
        'metiers': {
            'Agent d\'entretien': ['agent entretien', 'agent de propreté'],
            'Femme de ménage': ['femme de ménage', 'ménage'],
            'Laveur': ['laveur', 'lavage'],
        },
        'mots_cles': ['nettoyage', 'entretien', 'propreté', 'ménage']
    },
    'Immobilier': {
        'emoji': '🏠', 'couleur': '#E63946',
        'metiers': {
            'Agent immobilier': ['agent immobilier', 'immobilier'],
            'Promoteur': ['promoteur', 'promotion immobilière'],
            'Gestionnaire locatif': ['gestionnaire locatif', 'gestion locative'],
        },
        'mots_cles': ['immobilier', 'vente', 'location', 'bien']
    },
    'Sport et Loisirs': {
        'emoji': '⚽', 'couleur': '#FF6B00',
        'metiers': {
            'Éducateur sportif': ['éducateur sportif', 'entraîneur'],
            'Coach sportif': ['coach sportif'],
        },
        'mots_cles': ['sport', 'loisirs', 'animation', 'fitness']
    },
    'Non classé': {'emoji': '❓', 'couleur': '#888888', 'metiers': {}, 'mots_cles': []}
}

DOSSIER_TRI = "CV_classes"
DOSSIER_EXCEL = "exports"
DOSSIER_FILTRES = "CV_filtres"
DOSSIER_A_TRAITER = "cv_a_traiter"
for d in [DOSSIER_TRI, DOSSIER_EXCEL, DOSSIER_FILTRES, DOSSIER_A_TRAITER]:
    os.makedirs(d, exist_ok=True)

# ============================================================
# CONSTANTES APPRENTISSAGE
# ============================================================
FICHIER_APPRENTISSAGE = "apprentissage.json"
FICHIER_CONFIG_EMAIL = "config_email.json"
SEUIL_CONFIANCE_DEFAUT = 70

# ============================================================
# EXTRACTEUR
# ============================================================
class ExtracteurInfos:
    PRENOMS_FEMININS = {
        'marie', 'sophie', 'julie', 'claire', 'lucie', 'camille',
        'léa', 'chloé', 'céline', 'sandrine', 'nathalie', 'isabelle',
        'caroline', 'émilie', 'fatou', 'aïssatou', 'mariama',
        'aminata', 'kadiatou', 'awa', 'grace', 'espérance',
        'bénédicte', 'charlotte', 'sylvie', 'madeleine', 'jericae',
        'jeanne', 'françoise', 'monique', 'véronique', 'patricia',
        'catherine', 'christine', 'valérie', 'séverine', 'audrey',
        'jennifer', 'jessica', 'sarah', 'laura', 'emma', 'léna',
        'juliette', 'manon', 'élise', 'élodie', 'anaïs', 'océane',
        'jerica', 'ange', 'angele', 'angèle'
    }
    
    PRENOMS_MASCULINS = {
        'jean', 'pierre', 'paul', 'jacques', 'lucas', 'hugo',
        'antoine', 'nicolas', 'sébastien', 'thomas', 'julien',
        'maxime', 'alexandre', 'guillaume', 'françois', 'philippe',
        'laurent', 'moussa', 'ibrahima', 'amadou', 'oumar',
        'abdoulaye', 'mamadou', 'samuel', 'david', 'daniel',
        'emmanuel', 'joseph', 'alphonse', 'bernard', 'charles',
        'étienne', 'franck', 'gaston', 'henri', 'jules', 'kévin',
        'léon', 'michel', 'noël', 'oscar', 'pascal', 'raymond',
        'serge', 'victor', 'richard', 'sacha', 'marc', 'luc',
        'eric', 'éric', 'patrick', 'christophe', 'olivier',
        'vincent', 'benjamin', 'adrien', 'arthur', 'louis'
    }
    
    MOTS_INTERDITS = {
        'EXPÉRIENCES', 'EXPERIENCES', 'EXPÉRIENCE', 'EXPERIENCE',
        'FORMATION', 'FORMATIONS', 'ÉDUCATION', 'EDUCATION',
        'CERTIFICATIONS', 'CERTIFICATION', 'DIPLOMES', 'DIPLÔMES',
        'COMPÉTENCES', 'COMPETENCES', 'COMPÉTENCE', 'COMPETENCE',
        'LANGUES', 'LANGUE', 'INTÉRÊTS', 'INTERETS', 'INTÉRÊT',
        'CENTRES', 'CENTRE', 'LOISIRS', 'PROFIL', 'OBJECTIF',
        'OBJECTIFS', 'CONTACT', 'CONTACTS', 'INFORMATIONS',
        'RÉFÉRENCES', 'REFERENCES', 'PROJETS', 'PROJET',
        'ÉCOLE', 'ECOLE', 'ÉCOLES', 'UNIVERSITÉ', 'UNIVERSITE',
        'LYCÉE', 'LYCEE', 'COLLÈGE', 'COLLEGE', 'INSTITUT',
        'ENTREPRISE', 'ENTREPRISES', 'SOCIÉTÉ', 'SOCIETE',
        'STAGE', 'STAGES', 'EMPLOI', 'EMPLOIS', 'POSTE', 'POSTES',
        'CV', 'CURRICULUM', 'VITAE', 'RESUME', 'RÉSUMÉ',
        'NOM', 'PRENOM', 'PRÉNOM', 'NOMS', 'PRENOMS',
        'TEL', 'TÉL', 'TELEPHONE', 'TÉLÉPHONE', 'MOBILE',
        'EMAIL', 'MAIL', 'ADRESSE', 'VILLE', 'PAYS',
        'JANVIER', 'FÉVRIER', 'MARS', 'AVRIL', 'MAI', 'JUIN',
        'JUILLET', 'AOÛT', 'SEPTEMBRE', 'OCTOBRE', 'NOVEMBRE', 'DÉCEMBRE',
        'COMMUNICATION', 'INNOVATION', 'PRODUIT',
        'ANNÉE', 'PREMIÈRE', 'DEUXIÈME', 'TROISIÈME',
        'TCF', 'TOEIC', 'TOEFL', 'MARKETING', 'DIGITAL'
    }
    
    def __init__(self):
        self.villes = [
            'yaoundé', 'douala', 'bafoussam', 'garoua', 'bamenda',
            'maroua', 'ngaoundéré', 'bertoua', 'ebolowa',
            'kribi', 'limbe', 'buea', 'kumba', 'edea', 'dschang', 'foumban',
            'paris', 'lyon', 'marseille', 'toulouse', 'nice', 'nantes',
            'abidjan', 'dakar', 'lomé', 'cotonou', 'libreville',
            'brazzaville', 'kinshasa', 'casablanca', 'tunis', 'londres',
            'montréal', 'bruxelles'
        ]
    
    def extraire_email(self, texte):
        for p in [r'[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}',
                  r'[a-zA-Z0-9._%+-]+\s*@\s*[a-zA-Z0-9.-]+\s*\.\s*[a-zA-Z]{2,}']:
            m = re.search(p, texte, re.IGNORECASE)
            if m:
                return re.sub(r'\s+', '', m.group(0)).lower()
        for tld in ['com', 'fr', 'net', 'org', 'edu', 'cm', 'uk', 'de', 'ca']:
            pattern = r'[a-zA-Z0-9._%+-]+@[a-zA-Z0-9-]+' + tld + r'\b'
            m = re.search(pattern, texte, re.IGNORECASE)
            if m:
                email = m.group(0)
                return (email[:-len(tld)] + '.' + tld).lower()
        return "RAS"
    
    def extraire_tel(self, texte):
        for p in [r'\+237[\s.-]?\d[\s.-]?\d{2}[\s.-]?\d{2}[\s.-]?\d{2}[\s.-]?\d{2}',
                  r'00237[\s.-]?\d[\s.-]?\d{2}[\s.-]?\d{2}[\s.-]?\d{2}[\s.-]?\d{2}',
                  r'\b6[\s.-]?\d{2}[\s.-]?\d{2}[\s.-]?\d{2}[\s.-]?\d{2}\b',
                  r'0[1-9][\s.-]?\d{2}[\s.-]?\d{2}[\s.-]?\d{2}[\s.-]?\d{2}',
                  r'\+33[\s.-]?\d[\s.-]?\d{2}[\s.-]?\d{2}[\s.-]?\d{2}[\s.-]?\d{2}',
                  r'\+\d{1,3}[\s.-]?\(?\d{1,4}\)?[\s.-]?\d{2,4}[\s.-]?\d{2,4}[\s.-]?\d{0,4}',
                  r'\b\d{8,15}\b']:
            for m in re.findall(p, texte):
                tel = re.sub(r'[^\d+]', '', m)
                chiffres = re.sub(r'\D', '', tel)
                if len(chiffres) == 9 and chiffres.startswith('6'):
                    return f"+237 {chiffres[0]} {chiffres[1:3]} {chiffres[3:5]} {chiffres[5:7]} {chiffres[7:9]}"
                if 8 <= len(chiffres) <= 15:
                    return tel
        return "RAS"
    
    def extraire_nom_prenom(self, texte, email="RAS"):
        nom, prenom = "", ""
        for p in [r'(?:^|\n)\s*NOM\s*(?:DE\s*FAMILLE)?\s*[:\-]\s*([A-ZÀ-Ý][a-zà-ÿA-ZÀ-Ý\s\-]{1,30})']:
            m = re.search(p, texte, re.MULTILINE)
            if m:
                nom = m.group(1).strip().split()[0]
                break
        for p in [r'(?:^|\n)\s*PR[EÉ]NOM[S]?\s*[:\-]\s*([A-ZÀ-Ý][a-zà-ÿA-ZÀ-Ý\s\-]{1,30})']:
            m = re.search(p, texte, re.MULTILINE)
            if m:
                prenom = m.group(1).strip().split()[0]
                break
        if nom and prenom:
            return nom, prenom
        
        lignes = [l.strip() for l in texte.split('\n') if l.strip()]
        for ligne in lignes[:3]:
            ligne_clean = re.sub(r'[·•\-/–—|]', ' ', ligne)
            ligne_clean = re.sub(r'\s+', ' ', ligne_clean).strip()
            if len(ligne_clean) < 4 or len(ligne_clean) > 45:
                continue
            mots = re.findall(r'\b[A-Za-zÀ-ÿ]{2,20}\b', ligne_clean)
            if not (2 <= len(mots) <= 3):
                continue
            if any(m.upper() in self.MOTS_INTERDITS for m in mots):
                continue
            if not all(m[0].isupper() for m in mots):
                continue
            if len(mots) == 2:
                prenom = mots[0].capitalize()
                nom = mots[1].upper()
            elif len(mots) == 3:
                if mots[0].isupper() and mots[1].isupper() and mots[2].isupper():
                    nom = mots[0].upper()
                    prenom = f"{mots[1].capitalize()} {mots[2].capitalize()}"
                else:
                    prenom = f"{mots[0].capitalize()} {mots[1].capitalize()}"
                    nom = mots[2].upper()
            return nom, prenom
        
        if email and email != "RAS" and '@' in email:
            partie = re.sub(r'\d+', '', email.split('@')[0])
            for sep in ['.', '_', '-']:
                if sep in partie:
                    parts = [p for p in partie.split(sep) if p and p.upper() not in self.MOTS_INTERDITS and len(p) > 1]
                    if len(parts) >= 2:
                        return parts[1].upper(), parts[0].capitalize()
        return "RAS", "RAS"
    
    def extraire_niveau(self, texte):
        t = texte.lower()
        meilleur, pts = "RAS", 0
        for m in re.finditer(r'\bbac\s*\+\s*(\d+)\b', t):
            n = int(m.group(1))
            if n >= 8: niveau, p = 'Doctorat', 30
            elif n >= 5: niveau, p = 'Master/Ingénieur', 25
            elif n >= 3: niveau, p = 'Licence', 15
            elif n >= 2: niveau, p = 'BTS/DUT', 10
            else: niveau, p = 'Bac', 8
            if p > pts: pts, meilleur = p, niveau
        patterns = {
            r'\b(doctorat|phd|thèse|these)\b': ('Doctorat', 30),
            r'\b(master|mastère|msc|mba|dess|dea|ingénieur)\b': ('Master/Ingénieur', 25),
            r'\b(licence|bachelor|licencié)\b': ('Licence', 15),
            r'\b(bts|dut|deug|hnd)\b': ('BTS/DUT', 10),
            r'\b(baccalauréat|bac|gce)\b': ('Bac', 8),
        }
        for p, (niv, points) in patterns.items():
            if re.search(p, t, re.IGNORECASE) and points > pts:
                pts, meilleur = points, niv
        return meilleur, pts
    
    # Titres de section : on ignore les dates de la partie formation
    TITRE_FORMATION = re.compile(
        r"^\W*(formations?|[ée]ducation|[ée]tudes|dipl[ôo]mes?|cursus|"
        r"parcours\s+(?:acad[ée]mique|scolaire|universitaire))\b", re.IGNORECASE)
    TITRE_EXPERIENCE = re.compile(
        r"^\W*(exp[ée]riences?|parcours\s+professionnel|emplois?|stages?)\b", re.IGNORECASE)
    TITRE_AUTRE = re.compile(
        r"^\W*(comp[ée]tences|langues|certifications?|loisirs|centres?\s+d'int[ée]r[êe]ts?|"
        r"r[ée]f[ée]rences|projets|profil)\b", re.IGNORECASE)
    MOTS_ETUDES = re.compile(
        r"\b(master|licence|bachelor|baccalaur[ée]at|bac|bts|dut|doctorat|universit[ée]|[ée]cole|"
        r"lyc[ée]e|coll[èe]ge|dipl[ôo]me|[ée]tudiant|gce|mba)\b", re.IGNORECASE)
    MOTS_PRESENT = r"aujourd'hui|aujourd’hui|ce\s+jour|pr[ée]sent|present|maintenant|en\s+cours|now"
    PLAGE_DATES = re.compile(
        r"(?:\d{1,2}[/.]\s*)?(19\d{2}|20\d{2})\s*(?:-|–|—|\bà\b|\ba\b|\bau\b)\s*"
        r"(?:\d{1,2}[/.]\s*)?(19\d{2}|20\d{2}|" + MOTS_PRESENT + r")", re.IGNORECASE)
    DEPUIS = re.compile(r"\bdepuis\s+(?:\w+\s+)?(19\d{2}|20\d{2})", re.IGNORECASE)

    def extraire_experience(self, texte):
        """
        Années d'expérience professionnelle :
        1. "X ans d'expérience" écrit dans le CV → valeur directe
        2. Sinon, somme des périodes de travail (2018-2023, depuis 2019…),
           sans compter la partie formation et en fusionnant les périodes qui se chevauchent
        """
        annee_actuelle = datetime.now().year

        # 1. Valeur explicite
        explicites = []
        for p in [r"(\d{1,2})\s*\+?\s*(?:ans?|années?)\s*(?:d['’]\s*|de\s+)?(?:expérience|exp\b)",
                  r"expérience\s*(?:de|d['’])\s*(\d{1,2})\s*(?:ans?|années?)"]:
            for m in re.finditer(p, texte, re.IGNORECASE):
                v = int(m.group(1))
                if 0 < v <= 50:
                    explicites.append(v)
        if explicites:
            return max(explicites)

        # 2. Périodes de travail (hors section formation et lignes de diplômes)
        periodes = []
        section = None  # 'formation', 'experience' ou 'autre'
        titres = [('formation', self.TITRE_FORMATION), ('experience', self.TITRE_EXPERIENCE),
                  ('autre', self.TITRE_AUTRE)]
        for ligne in texte.split('\n'):
            if len(ligne.strip()) <= 40:
                titre = next((nom for nom, motif in titres if motif.match(ligne)), None)
                if titre:
                    section = titre
                    continue
            if section == 'formation':
                continue
            # Hors section "Expériences", une ligne qui parle de diplôme/école est une formation
            if section != 'experience' and self.MOTS_ETUDES.search(ligne):
                continue
            for debut, fin in self.PLAGE_DATES.findall(ligne):
                d = int(debut)
                f = int(fin) if fin.isdigit() else annee_actuelle
                if 1950 <= d <= f <= annee_actuelle + 1:
                    periodes.append((d, max(f, d)))
            for m in self.DEPUIS.finditer(ligne):
                d = int(m.group(1))
                if 1950 <= d <= annee_actuelle:
                    periodes.append((d, annee_actuelle))

        if not periodes:
            return 0
        # Fusion des périodes qui se chevauchent puis somme
        periodes.sort()
        total, (cour_d, cour_f) = 0, periodes[0]
        for d, f in periodes[1:]:
            if d <= cour_f:
                cour_f = max(cour_f, f)
            else:
                total += cour_f - cour_d
                cour_d, cour_f = d, f
        total += cour_f - cour_d
        return min(total, 50)

    def extraire_age(self, texte):
        """Âge du candidat (RAS si introuvable : on n'invente pas)"""
        return calculer_age_depuis_annee(texte)[0]

    def extraire_sexe(self, texte, prenom="RAS"):
        """Sexe du candidat : civilité > prénom connu > indices grammaticaux > RAS"""
        return determiner_sexe_avec_confiance(texte, prenom)[0]

    def extraire_situation(self, texte):
        t = texte.lower()
        # "mariée?" exige l'accent pour ne pas confondre avec le prénom "Marie"
        if re.search(r'\b(mariée?|mariage)\b', t): return "Marié(e)"
        if re.search(r'\b(c[ée]libataire|c[ée]libat)\b', t): return "Célibataire"
        if re.search(r'\b(divorc[ée]e?)\b', t): return "Divorcé(e)"
        if re.search(r'\b(veuf|veuve)\b', t): return "Veuf(ve)"
        return "RAS"
    
    def extraire_nationalite(self, texte):
        t = texte.lower()
        for cle, val in [('camerounais', 'Camerounaise'), ('camerounaise', 'Camerounaise'),
                         ('cameroun', 'Camerounaise'), ('français', 'Française'),
                         ('française', 'Française'), ('sénégalais', 'Sénégalaise'),
                         ('sénégalaise', 'Sénégalaise'), ('ivoirien', 'Ivoirienne'),
                         ('ivoirienne', 'Ivoirienne'), ('gabonais', 'Gabonaise'),
                         ('gabonaise', 'Gabonaise'), ('congolais', 'Congolaise'),
                         ('congolaise', 'Congolaise'), ('tchadien', 'Tchadienne'),
                         ('tchadienne', 'Tchadienne')]:
            if re.search(r'\b' + re.escape(cle) + r'\b', t):
                return val
        return "RAS"
    
    def extraire_langues(self, texte):
        t = texte.lower()
        trouvees = []
        for cle, nom in [('français', 'Français'), ('francais', 'Français'), ('french', 'Français'),
                         ('anglais', 'Anglais'), ('english', 'Anglais'),
                         ('espagnol', 'Espagnol'), ('spanish', 'Espagnol'),
                         ('allemand', 'Allemand'), ('italien', 'Italien'),
                         ('chinois', 'Chinois'), ('arabe', 'Arabe'),
                         ('portugais', 'Portugais'), ('russe', 'Russe')]:
            if re.search(r'\b' + re.escape(cle) + r'\b', t):
                if nom not in trouvees: trouvees.append(nom)
        return ', '.join(trouvees) if trouvees else "RAS"
    
    def extraire_competences(self, texte):
        t = texte.lower()
        trouvees = set()
        for data in CATEGORIES.values():
            for mot in data.get('mots_cles', []):
                if re.search(r'\b' + re.escape(mot) + r'\b', t):
                    trouvees.add(mot.capitalize())
        return sorted(trouvees)[:20]
    
    def extraire_certifications(self, texte):
        t = texte.lower()
        trouvees = []
        for c in ['pmp', 'prince2', 'itil', 'scrum master', 'aws', 'azure',
                  'google cloud', 'cisco', 'ccna', 'ccnp', 'microsoft',
                  'oracle', 'sap', 'comptia', 'ceh', 'google analytics',
                  'hubspot', 'toeic', 'toefl', 'ielts', 'tcf', 'delf',
                  'cpa', 'acca', 'cfa', 'cia', 'cisa', 'excel']:
            if re.search(r'\b' + re.escape(c) + r'\b', t):
                trouvees.append(c.upper())
        return ', '.join(trouvees) if trouvees else "RAS"
    
    def extraire_etablissement(self, texte):
        etabs = []
        for p in [r'universit[ée]\s+(?:de\s+|d\')?[A-ZÀ-Ý][a-zà-ÿ]+(?:\s+[A-ZÀ-Ý][a-zà-ÿ]+)?',
                  r'[ée]cole\s+(?:de\s+|des\s+|nationale\s+|sup[ée]rieure\s+)?[A-ZÀ-Ý][a-zà-ÿ]+(?:\s+[A-ZÀ-Ý][a-zà-ÿ]+)?',
                  r'institut\s+[A-ZÀ-Ý][a-zà-ÿ]+(?:\s+[A-ZÀ-Ý][a-zà-ÿ]+)?',
                  r'[A-ZÀ-Ý][a-zà-ÿ]+\s+College', r'Digital\s+College']:
            for m in re.finditer(p, texte):
                etab = m.group(0).strip()
                if etab not in etabs and 5 < len(etab) < 60:
                    etabs.append(etab)
        return ', '.join(etabs[:3]) if etabs else "RAS"
    
    def extraire_poste_actuel(self, texte):
        t = texte.lower()
        for poste in ['directeur', 'directrice', 'manager', 'responsable',
                      'chef de projet', 'ingénieur', 'consultant', 'analyste',
                      'assistant', 'assistante', 'technicien', 'comptable',
                      'développeur', 'commercial', 'secrétaire', 'stagiaire']:
            if re.search(r'\b' + re.escape(poste) + r'\b', t):
                return poste.capitalize()
        return "RAS"
    
    def extraire_ville(self, texte):
        t = texte.lower()
        for v in self.villes:
            if re.search(r'\b' + re.escape(v) + r'\b', t):
                return v.capitalize()
        return "RAS"
    
    def extraire_tout(self, texte):
        email = self.extraire_email(texte)
        nom, prenom = self.extraire_nom_prenom(texte, email)
        niveau, pts = self.extraire_niveau(texte)
        return {
            'nom': nom or "RAS", 'prenom': prenom or "RAS",
            'email': email, 'telephone': self.extraire_tel(texte),
            'age': self.extraire_age(texte),
            'sexe': self.extraire_sexe(texte, prenom),
            'situation': self.extraire_situation(texte),
            'nationalite': self.extraire_nationalite(texte),
            'niveau_etudes': niveau, 'points_niveau': pts,
            'etablissement': self.extraire_etablissement(texte),
            'annees_experience': self.extraire_experience(texte),
            'poste_actuel': self.extraire_poste_actuel(texte),
            'competences': self.extraire_competences(texte),
            'langues': self.extraire_langues(texte),
            'certifications': self.extraire_certifications(texte),
            'ville': self.extraire_ville(texte),
            'texte_brut': texte
        }

# ============================================================
# MODULE : ÂGE CALCULÉ + SEXE CONFIANCE
# ============================================================
# Une date n'est une date de naissance que si elle suit "né(e) le", "date de naissance"…
# (sinon les dates d'expérience comme "01/09/2018" donnaient des âges absurdes)
_CONTEXTE_NAISSANCE = (r"(?:(?<!\w)n[ée]\(?e?\)?\s+(?:le\s+|en\s+|à\s+\S+\s+le\s+)?"
                       r"|date\s+de\s+naissance\s*[:\-]?\s*|naissance\s*[:\-]\s*)")
_DATE_NAISSANCE = re.compile(
    _CONTEXTE_NAISSANCE + r"(?:\d{1,2}\s*[/\-.]\s*\d{1,2}\s*[/\-.]\s*|\d{1,2}(?:er)?\s+[a-zéû]+\s+)?(19\d{2}|20\d{2})",
    re.IGNORECASE)
_AGES_DIRECTS = [
    (re.compile(r"j['’]?ai\s+(\d{1,2})\s*ans", re.IGNORECASE), 0.90),
    (re.compile(r"(?<!\w)[âa]ge\s*[:\-]\s*(\d{1,2})\b", re.IGNORECASE), 0.95),  # pas "stage : 6 mois"
    (re.compile(r"\((\d{1,2})\s*ans\)", re.IGNORECASE), 0.80),
]


def calculer_age_depuis_annee(texte):
    """Âge du candidat et confiance : ("28", 0.98) ou ("RAS", 0.0)"""
    annee_actuelle = datetime.now().year
    ages_candidats = []
    for m in _DATE_NAISSANCE.finditer(texte):
        age = annee_actuelle - int(m.group(1))
        if 15 <= age <= 85: ages_candidats.append((age, 0.98))
    for motif, confiance in _AGES_DIRECTS:
        for m in motif.finditer(texte):
            age = int(m.group(1))
            if 15 <= age <= 85: ages_candidats.append((age, confiance))
    if not ages_candidats: return "RAS", 0.0
    ages_candidats.sort(key=lambda x: x[1], reverse=True)
    return str(ages_candidats[0][0]), ages_candidats[0][1]


def determiner_sexe_avec_confiance(texte, prenom="RAS"):
    """Sexe et confiance. Les pronoms "il"/"elle" ne sont pas utilisés (trop fréquents dans un texte)."""
    t = texte.lower()
    # Civilité dans l'en-tête du CV seulement (pas dans les références : "Contact : Mr X")
    entete = " ".join([l for l in t.split('\n') if l.strip()][:8])
    if re.search(r'\b(madame|mme|mlle|mademoiselle)\b', entete): return "Femme", 0.99
    if re.search(r'\b(monsieur|mr)\b', entete): return "Homme", 0.99
    if prenom and prenom != "RAS":
        for p in prenom.lower().split():
            if p in ExtracteurInfos.PRENOMS_FEMININS: return "Femme", 0.90
            if p in ExtracteurInfos.PRENOMS_MASCULINS: return "Homme", 0.90
    # Accords grammaticaux : "née le", "mariée" (mais pas la formule neutre "marié(e)")
    if re.search(r'(?<!\w)née\s+(?:le|en|à)\b', t): return "Femme", 0.90
    if re.search(r'(?<!\w)né\s+(?:le|en|à)\b', t): return "Homme", 0.90
    if re.search(r'\b(mariée|épouse|m[èe]re\s+de)\b', t): return "Femme", 0.85
    if re.search(r'\b(époux|p[èe]re\s+de)\b|\bmarié\b(?!\s*\(e\))', t): return "Homme", 0.85
    return "RAS", 0.0


def calculer_confiance_champ(valeur, champ, texte=""):
    if not valeur or valeur == "RAS": return 0.0
    if champ == 'email': return 0.98 if '@' in valeur and '.' in valeur else 0.5
    if champ == 'telephone':
        chiffres = re.sub(r'\D', '', valeur)
        return 0.95 if 8 <= len(chiffres) <= 15 else 0.5
    if champ in ['nom', 'prenom']: return 0.85 if len(valeur) >= 2 else 0.4
    if champ == 'niveau_etudes':
        return 0.90 if valeur in ['Bac', 'BTS/DUT', 'Licence', 'Master/Ingénieur', 'Doctorat'] else 0.4
    if champ == 'ville': return 0.85
    if champ == 'age': return 0.90
    if champ == 'sexe': return 0.85
    return 0.5


def verifier_champs_a_valider(infos, domaine, metier, seuil):
    champs_a_valider = []
    seuils = {
        'nom': seuil, 'prenom': seuil, 'email': 80, 'telephone': 80,
        'sexe': 80, 'age': 75, 'niveau_etudes': seuil, 'ville': seuil,
    }
    for champ, s in seuils.items():
        valeur = infos.get(champ, 'RAS')
        confiance = calculer_confiance_champ(valeur, champ)
        if confiance * 100 < s:
            champs_a_valider.append({
                'champ': champ, 'valeur_actuelle': valeur,
                'confiance': confiance, 'seuil': s
            })
    return champs_a_valider

# ============================================================
# MODULE : APPRENTISSAGE
# ============================================================
def apprentissage_vide():
    return {'sexe': {}, 'domaines': {}, 'metiers': {},
            'decisions': [],   # domaines validés par un humain : prédit → final
            'exemples': [],    # textes de CV étiquetés, utilisés pour réentraîner le modèle
            'options': {'auto': True, 'nouveau_domaine': False},
            'stats': {'total_corrections': 0, 'derniere_maj': None, 'version_modele': 0}}


def charger_apprentissage():
    app = apprentissage_vide()
    if os.path.exists(FICHIER_APPRENTISSAGE):
        try:
            with open(FICHIER_APPRENTISSAGE, 'r', encoding='utf-8') as f:
                donnees = json.load(f)
            # Fusion avec les valeurs par défaut (anciens fichiers sans décisions/exemples)
            for cle, val in donnees.items():
                if isinstance(val, dict) and isinstance(app.get(cle), dict):
                    app[cle].update(val)
                else:
                    app[cle] = val
        except: pass
    return app


def sauvegarder_apprentissage(app):
    try:
        app['stats']['derniere_maj'] = datetime.now().isoformat()
        with open(FICHIER_APPRENTISSAGE, 'w', encoding='utf-8') as f:
            json.dump(app, f, indent=2, ensure_ascii=False)
        return True
    except: return False


# ============================================================
# MODULE : CONFIGURATION MESSAGERIE (IMAP)
# ============================================================
def config_email_vide():
    return {
        'imap_server': '', 'imap_port': 993,
        'email_address': '', 'app_password': '',
        'dossier_traites': 'CV traités',
    }


def charger_config_email():
    cfg = config_email_vide()
    if os.path.exists(FICHIER_CONFIG_EMAIL):
        try:
            with open(FICHIER_CONFIG_EMAIL, 'r', encoding='utf-8') as f:
                donnees = json.load(f)
            for cle, val in donnees.items():
                if cle in cfg:
                    cfg[cle] = val
        except: pass
    return cfg


def sauvegarder_config_email(cfg):
    try:
        with open(FICHIER_CONFIG_EMAIL, 'w', encoding='utf-8') as f:
            json.dump(cfg, f, indent=2, ensure_ascii=False)
        return True
    except: return False


def creds_messagerie():
    """Identifiants IMAP : priorité à la configuration UI, sinon secrets.toml."""
    cfg = st.session_state.get('config_email') or config_email_vide()
    if cfg.get('imap_server') and cfg.get('email_address'):
        return (cfg['imap_server'], int(cfg.get('imap_port') or 993), cfg['email_address'],
                cfg.get('app_password', ''), cfg.get('dossier_traites', 'CV traités'))
    try:
        config = st.secrets["email"]
        return (config["imap_server"], int(config["imap_port"]), config["email_address"],
                config["app_password"], config.get("dossier_traites", "CV traités"))
    except Exception:
        return None


def appliquer_apprentissage(infos, app):
    if infos.get('sexe') == "RAS":
        prenom = infos.get('prenom', '').lower()
        if prenom in app.get('sexe', {}):
            infos['sexe'] = app['sexe'][prenom]
    return infos


def enregistrer_correction(app, champ, ancienne, nouvelle, prenom="RAS", sauvegarder=True):
    if not nouvelle or nouvelle == ancienne: return app
    # On apprend l'association prénom → sexe (clé = prénom, pas l'ancienne valeur)
    if champ == 'sexe' and nouvelle != 'RAS' and prenom and prenom != 'RAS':
        app.setdefault('sexe', {})[prenom.lower().strip()] = nouvelle
    app['stats']['total_corrections'] = app['stats'].get('total_corrections', 0) + 1
    if sauvegarder:
        sauvegarder_apprentissage(app)
    return app


def enregistrer_decision(app, r, final, apprendre=True):
    """Mémorise le domaine choisi par un humain pour un CV (et l'ajoute aux exemples d'entraînement)"""
    predit = r.get('categorie_predite', r.get('categorie', 'Non classé'))
    app['decisions'].append({'date': datetime.now().isoformat(timespec='seconds'),
                             'fichier': r.get('fichier', ''), 'predit': predit, 'final': final})
    if predit != final:
        app['stats']['total_corrections'] = app['stats'].get('total_corrections', 0) + 1
    texte = r.get('infos', {}).get('texte_brut', '')[:4000]
    if apprendre and final != 'Non classé' and texte.strip():
        app['exemples'].append({'texte': texte, 'domaine': final})
    sauvegarder_apprentissage(app)


def stats_decisions(app):
    """Précision du modèle mesurée sur les décisions humaines"""
    decisions = app.get('decisions', [])
    aujourd_hui = datetime.now().date().isoformat()
    corrections = [d for d in decisions if d['predit'] != d['final']]
    par_domaine = {}
    for d in decisions:
        if d['predit'] == 'Non classé': continue
        ok, n = par_domaine.get(d['predit'], (0, 0))
        par_domaine[d['predit']] = (ok + (d['predit'] == d['final']), n + 1)
    confusions = {}
    for d in corrections:
        if d['predit'] != 'Non classé':
            paire = (d['final'], d['predit'])
            confusions[paire] = confusions.get(paire, 0) + 1
    return {
        'nb_decisions': len(decisions),
        'precision': (100 * (len(decisions) - len(corrections)) / len(decisions)) if decisions else None,
        'corrections': len(corrections),
        'corrections_jour': sum(1 for d in corrections if d['date'].startswith(aujourd_hui)),
        'par_domaine': sorted(((dom, 100 * ok / n, n) for dom, (ok, n) in par_domaine.items()),
                              key=lambda x: x[1], reverse=True),
        'confusion': max(confusions, key=confusions.get) if confusions else None,
        'recentes': corrections[-5:][::-1],
    }


# ============================================================
# CLASSIFICATEUR INTELLIGENT
# ============================================================
class Classificateur:
    def __init__(self):
        self.pipeline = None
        self.exemples = []
        self.matrice_exemples = None
    
    def entrainer(self, exemples=None):
        textes, labels = [], []
        for cat, data in CATEGORIES.items():
            if cat == 'Non classé': continue
            for mot in data.get('mots_cles', []):
                textes.append(f"expérience {mot} compétences {mot} maîtrise {mot}")
                textes.append(f"spécialiste {mot} expert {mot}")
                labels.extend([cat, cat])
            for metier, mots in data.get('metiers', {}).items():
                for mot in mots:
                    textes.append(f"{mot} {cat} {metier}")
                    labels.append(cat)
        self.pipeline = Pipeline([
            ('tfidf', TfidfVectorizer(max_features=15000, stop_words=None,
                                      ngram_range=(1, 2), min_df=1, sublinear_tf=True)),
            ('clf', MultinomialNB(alpha=0.05))
        ])
        # CV validés par un humain (comptent triple face aux mots-clés génériques)
        exemples = exemples or []
        for ex in exemples:
            textes.extend([ex['texte']] * 3)
            labels.extend([ex['domaine']] * 3)
        self.pipeline.fit(textes, labels)
        self.exemples = exemples
        self.matrice_exemples = (self.pipeline.named_steps['tfidf'].transform([ex['texte'] for ex in exemples])
                                 if exemples else None)
    
    def exemple_proche(self, texte, seuil=0.75):
        """Domaine d'un CV déjà validé quasi identique (même candidat, CV renvoyé…)"""
        if self.matrice_exemples is None: return None
        similarites = (self.matrice_exemples @ self.pipeline.named_steps['tfidf'].transform([texte]).T).toarray().ravel()
        meilleur = int(np.argmax(similarites))
        if similarites[meilleur] >= seuil:
            return self.exemples[meilleur]['domaine'], float(similarites[meilleur])
        return None
    
    def suggerer(self, texte, n=3):
        """Les n domaines les plus probables selon le modèle : [(domaine, probabilité)]"""
        if not self.pipeline: self.entrainer()
        try:
            probas = self.pipeline.predict_proba([re.sub(r'\s+', ' ', texte.lower())])[0]
        except Exception:
            return []
        classes = self.pipeline.classes_
        return [(classes[i], float(probas[i])) for i in np.argsort(probas)[::-1][:n]]
    
    def classifier(self, texte):
        if not self.pipeline: self.entrainer()
        t = texte.lower()
        if len(t) < 20: return 'Non classé', 'RAS', 0.0
        
        # 0. Un CV quasi identique a déjà été validé par un humain
        proche = self.exemple_proche(texte)
        if proche:
            return proche[0], 'RAS', proche[1]
        
        # 1. Chercher LE POSTE dans les 15 premières lignes, puis dans tout le texte
        lignes = [l.strip() for l in texte.split('\n') if l.strip()]
        zone_titre = " ".join(lignes[:15]).lower()
        for zone, confiance in [(zone_titre, 0.85), (t, 0.7)]:
            trouve = chercher_metier(zone)
            if trouve:
                return trouve[0], trouve[1], confiance

        try:
            t_clean = re.sub(r'\s+', ' ', t)
            pred = self.pipeline.predict([t_clean])[0]
            proba = max(self.pipeline.predict_proba([t_clean])[0])
            if proba < 0.15: return 'Non classé', 'RAS', proba
            return pred, 'RAS', proba
        except:
            return 'Non classé', 'RAS', 0.0


# Postes reconnus directement dans le CV : mot-clé → (domaine, métier).
# Les mots-clés sont cherchés en mots entiers, le plus long d'abord
# ("ingénieur génie civil" passe avant "ingénieur civil").
METIERS_PRIORITAIRES = {
    'logisticien': ('Logistique et Transport', 'Logisticien'),
    'magasinier': ('Logistique et Transport', 'Magasinier'),
    'gestionnaire de stock': ('Logistique et Transport', 'Gestionnaire de stock'),
    'transitaire': ('Logistique et Transport', 'Transitaire'),
    'chauffeur': ('Logistique et Transport', 'Chauffeur'),
    'chef de dépôt': ('Logistique et Transport', 'Chef de dépôt'),
    'directeur général adjoint': ('Direction Générale', 'Directeur Général Adjoint'),
    'directeur général': ('Direction Générale', 'Directeur Général'),
    'dga': ('Direction Générale', 'Directeur Général Adjoint'),
    'dg': ('Direction Générale', 'Directeur Général'),
    'développeur': ('Informatique', 'Développeur'),
    'ingénieur informatique': ('Informatique', 'Développeur'),
    'ingénieur logiciel': ('Informatique', 'Développeur'),
    'architecte logiciel': ('Informatique', 'Développeur'),
    'architecte cloud': ('Informatique', 'DevOps'),
    'administrateur réseau': ('Informatique', 'Administrateur réseau'),
    'administrateur système': ('Informatique', 'Administrateur réseau'),
    'devops': ('Informatique', 'DevOps'),
    'data scientist': ('Informatique', 'Data Scientist'),
    'comptable': ('Comptabilité', 'Comptable'),
    'auditeur': ('Comptabilité', 'Auditeur'),
    'expert-comptable': ('Comptabilité', 'Expert-comptable'),
    'analyste financier': ('Finance', 'Analyste financier'),
    'ressources humaines': ('Ressources Humaines', 'Chargé RH'),
    'recruteur': ('Ressources Humaines', 'Recruteur'),
    'drh': ('Ressources Humaines', 'DRH'),
    'marketing digital': ('Marketing', 'Marketing digital'),
    'community manager': ('Marketing', 'Marketing digital'),
    'commercial': ('Commercial et Vente', 'Commercial'),
    'médecin': ('Santé et Médical', 'Médecin'),
    'infirmier': ('Santé et Médical', 'Infirmier'),
    'infirmière': ('Santé et Médical', 'Infirmier'),
    'enseignant': ('Enseignement', 'Enseignant'),
    'formateur': ('Formation professionnelle', 'Formateur'),
    'maçon': ('BTP et Construction', 'Maçon'),
    # "ingénieur" seul est trop vague (il envoyait tout le monde en Génie Civil) :
    # on ne garde que les spécialités, le reste est laissé au modèle
    'ingénieur civil': ('Génie Civil', 'Ingénieur civil'),
    'ingénieur génie civil': ('Génie Civil', 'Ingénieur civil'),
    'ingénieur structure': ('Génie Civil', 'Ingénieur civil'),
    'ingénieur électrique': ('Génie Électrique', 'Ingénieur électrique'),
    'ingénieur électricité': ('Génie Électrique', 'Ingénieur électrique'),
    'ingénieur mécanique': ('Génie Mécanique', 'Ingénieur mécanique'),
    'ingénieur procédés': ('Génie des Procédés', 'Ingénieur procédés'),
    'ingénieur chimiste': ('Génie des Procédés', 'Ingénieur chimiste'),
    'ingénieur télécom': ('Télécommunications', 'Ingénieur télécom'),
    'ingénieur télécommunications': ('Télécommunications', 'Ingénieur télécom'),
    'ingénieur pétrole': ('Énergie et Pétrole', 'Ingénieur pétrole'),
    'ingénieur agronome': ('Agriculture', 'Agronome'),
    'qhse': ('QHSE', 'Responsable QHSE'),
    'mécanicien auto': ('Mécanique Auto', 'Mécanicien auto'),
    'mécanicien industriel': ('Génie Mécanique', 'Technicien mécanique'),
    'mécanicien': ('Mécanique Auto', 'Mécanicien auto'),
    'vidangeur': ('Mécanique Auto', 'Vidangeur'),
    'agent de sécurité': ('Sécurité', 'Agent de sécurité'),
    'cuisinier': ('Hôtellerie et Restauration', 'Cuisinier'),
    'serveur': ('Hôtellerie et Restauration', 'Serveur'),
    'serveuse': ('Hôtellerie et Restauration', 'Serveur'),
    'secrétaire': ('Administration', 'Secrétaire'),
    'juriste': ('Droit et Juridique', 'Juriste'),
    'architecte': ('Architecture et Urbanisme', 'Architecte'),
}

# Mot-clé ignoré si l'un de ces mots est autour ("serveur Windows", "centre commercial"…)
CONTEXTES_EXCLUS = {
    'serveur': ['windows', 'linux', 'web', 'réseau', 'informatique', 'administrat', 'cloud', 'sql', 'mail', 'dns'],
    'commercial': ['centre commercial', 'registre du commerce'],
    'architecte': ['logiciel', 'cloud', 'solution', 'données', 'système', 'réseau'],
}
NEGATIONS = ['pas de ', 'non ', 'aucun ']
_MOTIFS_METIERS = [(re.compile(r'(?<!\w)' + re.escape(mot) + r'(?!\w)'), mot, dom, met)
                   for mot, (dom, met) in sorted(METIERS_PRIORITAIRES.items(), key=lambda x: -len(x[0]))]


def chercher_metier(zone):
    """Métier connu trouvé en mot entier dans la zone (le mot-clé le plus long gagne)"""
    for motif, mot, domaine, metier in _MOTIFS_METIERS:
        for m in motif.finditer(zone):
            contexte = zone[max(0, m.start() - 30):m.end() + 30]
            if any(neg in contexte for neg in NEGATIONS + CONTEXTES_EXCLUS.get(mot, [])):
                continue
            return domaine, metier
    return None


# ============================================================
# SCORE
# ============================================================
def calculer_score(infos):
    pts_dip = min(infos.get('points_niveau', 0), 30)
    exp = infos.get('annees_experience', 0)
    if exp >= 10: pts_exp = 40
    elif exp >= 7: pts_exp = 35
    elif exp >= 5: pts_exp = 30
    elif exp >= 3: pts_exp = 22
    elif exp >= 2: pts_exp = 15
    elif exp >= 1: pts_exp = 8
    else: pts_exp = 0
    pts_comp = 0
    if infos.get('nom') and infos.get('nom') != "RAS": pts_comp += 3
    if infos.get('prenom') and infos.get('prenom') != "RAS": pts_comp += 3
    if infos.get('email') and infos.get('email') != "RAS": pts_comp += 8
    if infos.get('telephone') and infos.get('telephone') != "RAS": pts_comp += 8
    if infos.get('competences') and len(infos['competences']) >= 3: pts_comp += 4
    elif infos.get('competences'): pts_comp += 2
    if infos.get('ville') and infos.get('ville') != "RAS": pts_comp += 4
    pts_comp = min(pts_comp, 30)
    return {'total': pts_dip + pts_exp + pts_comp,
            'diplome': pts_dip, 'experience': pts_exp, 'completude': pts_comp}

# ============================================================
# LECTURE CV
# ============================================================
def _texte_ocr(resultat):
    """Normalise la sortie de RapidOCR (compatible tuple v1.2/v1.3 et objet v1.4+)."""
    if resultat is None:
        return ""
    if isinstance(resultat, tuple):
        resultat = resultat[0]
    if resultat is None:
        return ""
    if hasattr(resultat, 'txts'):
        return "\n".join(t for t in resultat.txts if t)
    return "\n".join(item[1] for item in resultat if len(item) > 1 and item[1])


def lire_avec_ocr(chemin):
    reader = charger_ocr()
    if not reader: return ""
    try:
        if chemin.lower().endswith('.pdf'):
            doc = fitz.open(chemin)
            texte_final = ""
            for page in doc:
                pix = page.get_pixmap(dpi=200)  # 200 dpi : bien plus lisible pour l'OCR
                img_bytes = pix.tobytes("png")
                texte_final += _texte_ocr(reader(img_bytes)) + "\n"
            doc.close()
            return texte_final
        else:
            return _texte_ocr(reader(chemin))
    except: return ""


def lire_cv(chemin):
    """Texte du CV. Le cache tient compte de la date et de la taille du fichier :
    un CV remplacé par un autre portant le même nom est relu."""
    try:
        infos_fichier = os.stat(chemin)
    except OSError:
        return ""
    return _lire_cv_cache(chemin, infos_fichier.st_mtime, infos_fichier.st_size)


@st.cache_data(show_spinner=False)
def _lire_cv_cache(chemin, date_modif, taille):
    # date_modif et taille ne servent qu'à la clé de cache
    texte = ""
    try:
        if chemin.lower().endswith('.pdf'):
            doc = fitz.open(chemin)
            for page in doc: texte += page.get_text() + "\n"
            doc.close()
        elif chemin.lower().endswith('.docx'):
            d = docx.Document(chemin)
            texte = "\n".join(p.text for p in d.paragraphs)
            for tab in d.tables:
                for r in tab.rows:
                    for c in r.cells: texte += " " + c.text
            return texte
    except: pass
    if len(texte.strip()) >= 300: return texte
    texte_ocr = lire_avec_ocr(chemin)
    if texte_ocr and len(texte_ocr.strip()) > len(texte.strip()):
        return texte_ocr
    return texte

# ============================================================
# EXPORT EXCEL
# ============================================================
def excel_octets(resultats):
    """Classeur Excel des CV analysés, en mémoire (None si aucun CV lisible)"""
    data = []
    for r in resultats:
        if not r.get('succes'): continue
        i = r['infos']
        s = r['scores']
        data.append({
            'Nom': i.get('nom', 'RAS'), 'Prénom': i.get('prenom', 'RAS'),
            'Email': i.get('email', 'RAS'), 'Téléphone': i.get('telephone', 'RAS'),
            'Âge': i.get('age', 'RAS'), 'Sexe': i.get('sexe', 'RAS'),
            'Situation': i.get('situation', 'RAS'),
            'Nationalité': i.get('nationalite', 'RAS'),
            'Ville': i.get('ville', 'RAS'),
            'Niveau': i.get('niveau_etudes', 'RAS'),
            'Établissement': i.get('etablissement', 'RAS'),
            'Expérience': i.get('annees_experience', 0),
            'Poste actuel': i.get('poste_actuel', 'RAS'),
            'Compétences': ', '.join(i.get('competences', [])),
            'Langues': i.get('langues', 'RAS'),
            'Certifications': i.get('certifications', 'RAS'),
            'Domaine': r['categorie'], 'Métier': r.get('metier', 'RAS'),
            'Confiance': f"{r['confiance']*100:.0f}%",
            'Score Total': s['total'], 'Score Diplôme': s['diplome'],
            'Score Expérience': s['experience'], 'Score Complétude': s['completude'],
            'Résumé': r.get('resume', '')
        })
    if not data: return None
    df = pd.DataFrame(data).sort_values('Score Total', ascending=False)
    df.insert(0, 'Rang', range(1, len(df) + 1))
    tampon = io.BytesIO()
    with pd.ExcelWriter(tampon, engine='openpyxl') as w:
        df.to_excel(w, sheet_name='CV', index=False)
        ws = w.sheets['CV']
        for col in ws.columns:
            m = max(len(str(c.value)) if c.value else 0 for c in col)
            ws.column_dimensions[col[0].column_letter].width = min(m + 2, 50)
    return tampon.getvalue()


def nom_export(prefixe="CVSmart_export"):
    return f"{prefixe}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"


# ============================================================
# MATCHING
# ============================================================
def calculer_matching(infos, offre):
    offre_l = offre.lower()
    cv_l = infos.get('texte_brut', '').lower()
    mots_offre = set(re.findall(r'\b[a-zà-ÿ]{3,}\b', offre_l))
    mots_offre -= {'avec', 'pour', 'dans', 'vous', 'nous', 'les', 'des', 'une',
                   'que', 'qui', 'sur', 'est', 'sont', 'être', 'avoir', 'plus'}
    mots_cv = set(re.findall(r'\b[a-zà-ÿ]{3,}\b', cv_l))
    communs = mots_offre & mots_cv
    score = (len(communs) / len(mots_offre) * 100) if mots_offre else 0
    m = re.search(r'(\d+)\s*(?:ans?|années?)\s*(?:d\')?expérience', offre_l)
    if m:
        exp_d = int(m.group(1))
        exp_cv = infos.get('annees_experience', 0)
        score = min(score + 10, 100) if exp_cv >= exp_d else max(score - 15, 0)
    return {'score': round(score, 1),
            'communs': sorted(communs)[:10],
            'manquants': sorted(mots_offre - mots_cv)[:5]}

# ============================================================
# RÉSUMÉ
# ============================================================
def generer_resume(infos, categorie, metier="RAS"):
    nom = f"{infos.get('prenom', '')} {infos.get('nom', '')}".strip()
    if nom in ["RAS RAS", "RAS", ""]: nom = "Ce candidat"
    niveau = infos.get('niveau_etudes', 'RAS')
    exp = infos.get('annees_experience', 0)
    comps = infos.get('competences', [])[:5]
    ville = infos.get('ville', 'RAS')
    p1 = f"{nom}"
    if metier and metier != "RAS": p1 += f", {metier}"
    elif niveau != "RAS": p1 += f", titulaire d'un {niveau}"
    if ville != "RAS": p1 += f", basé à {ville}"
    p1 += "."
    if exp > 0:
        p2 = f"Possède {exp} an{'s' if exp > 1 else ''} d'expérience"
        if categorie and categorie != 'Non classé':
            p2 += f" dans le domaine {categorie.lower()}"
        p2 += "."
    else:
        p2 = "Profil débutant, prêt à s'investir."
    p3 = f"Compétences clés : {', '.join(comps[:4])}." if comps else ""
    return f"{p1} {p2} {p3}".strip()

# ============================================================
# CHARGEMENT DES MODÈLES
# ============================================================
@st.cache_resource(show_spinner=False)
def charger_ocr():
    """Charge RapidOCR (ONNX, sans PyTorch) — uniquement à la demande, pour un CV scanné."""
    try:
        from rapidocr_onnxruntime import RapidOCR
        return RapidOCR()
    except: return None


def signature_exemples(exemples):
    """Empreinte des exemples d'entraînement : sert de clé de cache au modèle"""
    h = hashlib.sha1()
    for ex in exemples:
        h.update(ex.get('domaine', '').encode('utf-8') + b'\0' + ex.get('texte', '').encode('utf-8') + b'\0')
    return f"{len(exemples)}-{h.hexdigest()}"


@st.cache_resource(show_spinner=False)
def charger_classificateur(_exemples=(), signature=""):
    # _exemples n'est pas haché par Streamlit : la clé de cache est la signature des exemples,
    # donc un modèle en cache correspond toujours exactement aux exemples fournis
    classif = Classificateur()
    classif.entrainer(list(_exemples))
    return classif


def obtenir_classificateur(app):
    return charger_classificateur(app['exemples'], signature_exemples(app['exemples']))


def reentrainer_classificateur():
    app = st.session_state.apprentissage
    app['stats']['version_modele'] = app['stats'].get('version_modele', 0) + 1
    sauvegarder_apprentissage(app)
    st.session_state.classif = obtenir_classificateur(app)

# ============================================================
# SESSION ET NAVIGATION
# ============================================================
PAGES = [
    ("Accueil", "nav_accueil"),
    ("Analyser", "nav_analyser"),
    ("Emails", "nav_emails"),
    ("Résultats", "nav_resultats"),
    ("Filtrer", "nav_filtrer"),
    ("Validation", "nav_validation"),
    ("Apprentissage", "nav_apprentissage"),
    ("Correspondance", "nav_correspondance"),
    ("Paramètres", "nav_parametres"),
]


def init_session():
    if 'apprentissage' not in st.session_state:
        st.session_state.apprentissage = charger_apprentissage()
    if 'classif' not in st.session_state:
        app = st.session_state.apprentissage
        with st.spinner("Initialisation du modèle..."):
            st.session_state.classif = obtenir_classificateur(app)
    if 'extract' not in st.session_state:
        st.session_state.extract = ExtracteurInfos()
    if 'config_email' not in st.session_state:
        st.session_state.config_email = charger_config_email()
    defauts = {
        'resultats': [], 'cv_a_valider': [], 'cv_corriges': [], 'journal': [],
        'page_active': "Accueil", 'seuil_confiance': SEUIL_CONFIANCE_DEFAUT,
        'chemin_dossier': os.path.abspath(DOSSIER_A_TRAITER), 'edition_dossier': False,
        'derniere_analyse': None, 'validation_sel': None, 'flash': None,
    }
    for cle, val in defauts.items():
        if cle not in st.session_state:
            st.session_state[cle] = val


def aller_a(page):
    st.session_state.page_active = page


def flash(texte, ton='vert', icone_nom='valide'):
    """Message affiché en haut de la page au prochain affichage (survit à st.rerun)"""
    st.session_state.flash = (texte, ton, icone_nom)


def afficher_menu():
    with st.sidebar:
        ui.afficher(f'<div class="cv-marque">{ui.LOGO_MARQUE}</div>')
        for page, cle in PAGES:
            actif = st.session_state.page_active == page
            st.button(page, key=cle, type="primary" if actif else "secondary",
                      use_container_width=True, on_click=aller_a, args=(page,))
        ui.afficher(f'<div class="cv-pied">CVSmart v{VERSION_APP}<br>CIBLE RH Emploi</div>')

# ============================================================
# OUTILS COMMUNS AUX PAGES
# ============================================================
EXTENSIONS_CV = ('.pdf', '.docx', '.doc')


def lister_cv(dossier):
    if not os.path.isdir(dossier): return []
    return sorted(f for f in os.listdir(dossier) if f.lower().endswith(EXTENSIONS_CV))


def nom_dossier(texte):
    return texte.replace(" ", "_").replace("/", "-")


def nom_court(infos):
    """« Marie N. » comme sur les maquettes"""
    prenom, nom = infos.get('prenom', 'RAS'), infos.get('nom', 'RAS')
    if prenom != 'RAS' and nom != 'RAS': return f"{prenom} {nom[0].upper()}."
    if prenom != 'RAS': return prenom
    if nom != 'RAS': return nom.capitalize()
    return "Candidat inconnu"


def texte_experience(annees):
    if not annees: return "< 1 an"
    return f"{annees} an{'s' if annees > 1 else ''}"


def analyser_cv(chemin, nom, source=None, copie_nom=None, sous_dossier_metier=False):
    """Lit, extrait, classe un CV et l'ajoute aux résultats. Renvoie le résultat."""
    texte = lire_cv(chemin)
    if not texte or not texte.strip():
        r = {'fichier': nom, 'chemin': chemin, 'succes': False, 'erreur': 'Lecture impossible'}
        st.session_state.resultats.append(r)
        return r

    infos = st.session_state.extract.extraire_tout(texte)
    infos = appliquer_apprentissage(infos, st.session_state.apprentissage)

    cat, metier, conf = st.session_state.classif.classifier(texte)
    cat, metier, conf = str(cat), str(metier), float(conf)

    # Copie dans le dossier du domaine (et du métier pour les CV reçus par email)
    dest = os.path.join(DOSSIER_TRI, nom_dossier(cat))
    if sous_dossier_metier:
        dest = os.path.join(dest, nom_dossier(metier) if metier != "RAS" else "_Non_renseigné")
    os.makedirs(dest, exist_ok=True)
    copie = os.path.join(dest, copie_nom or nom)
    try: shutil.copy2(chemin, copie)
    except: copie = None

    r = {'fichier': nom, 'chemin': chemin, 'copie': copie, 'succes': True,
         'categorie': cat, 'categorie_predite': cat, 'metier': metier, 'confiance': conf,
         'infos': infos, 'scores': calculer_score(infos), 'resume': generer_resume(infos, cat, metier)}
    if source:
        r['source'] = source
    st.session_state.resultats.append(r)

    champs_faibles = verifier_champs_a_valider(infos, cat, metier, st.session_state.seuil_confiance)
    r['a_valider'] = bool(cat == 'Non classé' or conf < st.session_state.seuil_confiance / 100 or len(champs_faibles) >= 2)
    if r['a_valider']:
        st.session_state.cv_a_valider.append(len(st.session_state.resultats) - 1)
    return r


def changer_domaine(r, domaine):
    """Applique un nouveau domaine à un résultat et déplace sa copie classée"""
    if domaine == r['categorie']: return
    r['categorie'] = domaine
    if r.get('metier') not in CATEGORIES.get(domaine, {}).get('metiers', {}):
        r['metier'] = 'RAS'
    r['resume'] = generer_resume(r['infos'], domaine, r['metier'])
    if r.get('copie') and os.path.exists(r['copie']):
        dest = os.path.join(DOSSIER_TRI, nom_dossier(domaine))
        os.makedirs(dest, exist_ok=True)
        try:
            nouvelle = os.path.join(dest, os.path.basename(r['copie']))
            shutil.move(r['copie'], nouvelle)
            r['copie'] = nouvelle
        except: pass


def bouton_export(libelle, resultats, cle, prefixe="CVSmart_export"):
    """Bouton de téléchargement Excel (une copie est aussi enregistrée dans exports/)"""
    donnees = excel_octets(resultats)
    nom = nom_export(prefixe)

    def sauvegarder():
        with open(os.path.join(DOSSIER_EXCEL, nom), 'wb') as f:
            f.write(donnees)

    st.download_button(libelle, data=donnees or b"", file_name=nom, key=cle, type="primary",
                       mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                       disabled=not donnees, use_container_width=True,
                       on_click=sauvegarder if donnees else None)


def etat_vide(texte, libelle_bouton=None, page=None):
    ui.message(texte, 'bleu', 'boite')
    if libelle_bouton:
        st.button(libelle_bouton, key=f"vide_{page}", type="primary", on_click=aller_a, args=(page,))

# ============================================================
# PAGE ACCUEIL
# ============================================================
def ouvrir_selection_dossier():
    st.session_state.page_active = "Analyser"
    st.session_state.edition_dossier = True
    st.session_state.saisie_dossier = st.session_state.chemin_dossier


def page_accueil():
    B, R = ui.BLEU, ui.ROUGE
    ui.afficher('<div class="cv-coin"></div>')
    _, centre, _ = st.columns([1, 4.6, 1])
    with centre:
        ui.afficher(f'''<div class="cv-hero">
            <div class="cv-hero-logo">{ui.LOGO_CVSMART}<div class="cv-hero-nom"><span style="color:{B}">CV</span><span style="color:{R}">Smart</span></div></div>
            <div class="cv-hero-accroche">L'intelligence qui trie vos talents</div>
            <div class="cv-hero-texte">Classez automatiquement vos CV par domaine professionnel en quelques clics.</div>
        </div>''')
        _, b, _ = st.columns([1, 2, 1])
        with b:
            st.button("Sélectionner un dossier de CV", key="btn_accueil_dossier", type="primary",
                      use_container_width=True, on_click=ouvrir_selection_dossier)
        ui.afficher('<div class="cv-muted" style="text-align:center;margin-top:-0.4rem">Formats acceptés : PDF, DOCX, DOC</div>')

        atouts = [('eclair', 'Rapide', 'Analyse en quelques secondes', R),
                  ('cible', 'Précis', 'Classification par domaine', B),
                  ('bouclier', 'Simple', 'Interface intuitive', B)]
        ui.afficher('<div class="cv-atouts" style="margin-top:1.4rem">' + ''.join(
            f'<div>{ui.icone(ic, 24, coul)}<div class="cv-atout-titre" style="color:{coul}">{t}</div>'
            f'<div class="cv-muted">{d}</div></div>' for ic, t, d, coul in atouts) + '</div>')

        ui.afficher('<div class="cv-label" style="margin:1.3rem 0 -0.3rem 0">Nouveaux outils</div>')
        outils = [("Filtrer", "Affiner par domaine, expérience, compétences"),
                  ("Validation", "Confirmer ou corriger les classements"),
                  ("Apprentissage", "Améliorer le tri grâce à vos corrections"),
                  ("Correspondance", "Comparer les CV à une offre d'emploi")]
        for col, (page, texte) in zip(st.columns(4), outils):
            with col, st.container(key=f"clic_outil_{page.lower()}"):
                ui.afficher(f'<div class="cv-outil"><div class="cv-outil-titre">{page}</div>'
                            f'<div class="cv-outil-texte">{texte}</div></div>')
                st.button(f"Ouvrir {page}", key=f"outil_{page.lower()}", on_click=aller_a, args=(page,))

# ============================================================
# PAGE ANALYSER
# ============================================================
def basculer_edition_dossier():
    if st.session_state.edition_dossier:
        saisie = st.session_state.get('saisie_dossier', '').strip()
        if saisie:
            st.session_state.chemin_dossier = saisie
            st.session_state.derniere_analyse = None
    else:
        st.session_state.saisie_dossier = st.session_state.chemin_dossier
    st.session_state.edition_dossier = not st.session_state.edition_dossier


def rendu_progression(total, termines, journal, en_cours=False):
    restants = total - termines if en_cours else 0
    cartes = ui.grille([
        ui.carte_stat("Nombre de CV", total, icone_nom='fichier', ton='rose'),
        ui.carte_stat("En cours d'analyse", restants, icone_nom='horloge', ton='bleu'),
        ui.carte_stat("Terminés", termines, icone_nom='valide', ton='vert'),
    ], 3)
    pct = 100 * termines / total if total else 0
    progression = (f'<div class="cv-carte" style="margin-top:1.2rem"><div style="display:flex;justify-content:space-between">'
                   f'<span class="cv-h3" style="margin:0">Progression</span>'
                   f'<span class="cv-muted" style="font-weight:700">{termines} / {total}</span></div>'
                   f'<div class="cv-progression"><div style="width:{pct:.1f}%"></div></div></div>')
    if journal:
        lignes = ''.join(
            f'<div class="cv-ligne"><span class="cv-heure">{e["heure"]}</span><span class="cv-msg">{e["html"]}</span>'
            f'{"<span class=cv-spinner></span>" if e.get("en_cours") else ""}</div>'
            for e in reversed(journal[-60:]))
    else:
        lignes = ('<div class="cv-muted" style="margin:auto;text-align:center">Aucune analyse pour le moment. '
                  'Cliquez sur « Lancer l\'analyse » pour classer les CV du dossier.</div>')
    return f'{cartes}{progression}<div class="cv-carte cv-journal" style="margin-top:1.2rem">{lignes}</div>'


def page_analyser():
    ui.entete("Analyse et classement des CV")
    chemin_dossier = st.session_state.chemin_dossier

    with st.container(key="carte_dossier"):
        c1, c2, c3 = st.columns([5, 1.2, 1.7], vertical_alignment="center")
        with c1:
            if st.session_state.edition_dossier:
                st.text_input("Dossier sélectionné", key="saisie_dossier")
            else:
                ui.afficher(f'<div class="cv-h3" style="margin-bottom:0.2rem">Dossier sélectionné</div>'
                            f'<div style="font-family:ui-monospace,Menlo,monospace;color:#3A4357;word-break:break-all">'
                            f'{escape(chemin_dossier)}</div>')
        with c2:
            st.button("Valider" if st.session_state.edition_dossier else "Changer", key="btn_changer",
                      use_container_width=True, on_click=basculer_edition_dossier)
        with c3:
            lancer = st.button("Lancer l'analyse", type="primary", key="btn_lancer", use_container_width=True,
                               disabled=st.session_state.edition_dossier)

    fichiers = lister_cv(chemin_dossier)
    zone = st.empty()

    if lancer:
        if not os.path.isdir(chemin_dossier):
            ui.message(f"Dossier introuvable : {escape(chemin_dossier)}", 'rose', 'erreur')
            return
        if not fichiers:
            ui.message(f"Aucun CV (PDF, DOCX, DOC) dans {escape(chemin_dossier)}", 'ambre', 'alerte')
            return
        journal = st.session_state.journal = []
        total = len(fichiers)
        for i, nom in enumerate(fichiers):
            journal.append({'heure': datetime.now().strftime("%H:%M"), 'html': f"Analyse de {escape(nom)} ...", 'en_cours': True})
            zone.html(rendu_progression(total, i, journal, en_cours=True))
            r = analyser_cv(os.path.join(chemin_dossier, nom), nom)
            journal[-1]['en_cours'] = False
            heure = datetime.now().strftime("%H:%M")
            if not r['succes']:
                journal.append({'heure': heure, 'html': f'<span style="color:{ui.ROUGE}">Lecture impossible : {escape(nom)}</span>'})
                continue
            cat = escape(r['categorie'])
            suffixe = f' <span style="color:{ui.TONS["ambre"][1]}">· à valider</span>' if r['a_valider'] else ''
            journal.append({'heure': heure, 'html': f'Classification – <b style="color:{ui.VERT}">{cat}</b>{suffixe}'})
            if r.get('copie'):
                journal.append({'heure': heure, 'html': f'Copie vers le dossier <b>{cat}</b>'})
        nb_valider = sum(1 for r in st.session_state.resultats[-total:] if r.get('a_valider'))
        journal.append({'heure': datetime.now().strftime("%H:%M"),
                        'html': f'<b>Analyse terminée</b> – {total} CV traités'
                                + (f', <b style="color:{ui.TONS["ambre"][1]}">{nb_valider} à valider</b>' if nb_valider else '')})
        st.session_state.derniere_analyse = {'total': total, 'termines': total}

    derniere = st.session_state.derniere_analyse or {'total': len(fichiers), 'termines': 0}
    zone.html(rendu_progression(derniere['total'], derniere['termines'], st.session_state.journal))

# ============================================================
# PAGE RÉSULTATS
# ============================================================
def page_resultats():
    resultats = st.session_state.resultats
    c1, c2, c3 = st.columns([4.2, 1, 1.9], vertical_alignment="center")
    with c1: ui.entete("Résultats de classification")
    with c2: st.button("Filtrer", key="btn_res_filtrer", use_container_width=True, on_click=aller_a, args=("Filtrer",))
    with c3: bouton_export("Exporter les résultats", resultats, "btn_res_export")

    if not resultats:
        etat_vide("Aucun résultat pour le moment. Lancez une analyse pour classer vos CV.", "Analyser des CV", "Analyser")
        return

    classes = [r for r in resultats if r.get('succes') and r['categorie'] != 'Non classé']
    nb_non_classes = len(resultats) - len(classes)
    ui.afficher(ui.grille([
        ui.carte_stat("Total des CV", len(resultats)),
        ui.carte_stat("Classés", len(classes), ui.VERT),
        ui.carte_stat("Non classés", nb_non_classes, ui.ROUGE),
    ], 3))

    compteur = {}
    for r in classes:
        compteur[r['categorie']] = compteur.get(r['categorie'], 0) + 1
    tuiles = []
    for cat, nb in sorted(compteur.items(), key=lambda x: -x[1]):
        fond, texte = ui.ton_domaine(cat)
        tuiles.append(f'<div class="cv-carte" style="background:{fond};border:none"><div style="font-weight:700;color:{texte}">'
                      f'{escape(cat)}</div><div style="color:#3A4357;margin-top:0.3rem">{nb} CV</div></div>')
    if nb_non_classes:
        tuiles.append(f'<div class="cv-carte" style="background:#ECEFF4;border:1.5px dashed #9AA3B5">'
                      f'<div style="font-weight:700;color:#3A4357">Non classé</div>'
                      f'<div style="color:#3A4357;margin-top:0.3rem">{nb_non_classes} CV · à valider</div></div>')
    ui.afficher('<div class="cv-h2" style="margin-top:0.6rem">Répartition par domaine</div>' + ui.grille(tuiles, 4))

    nb = len(st.session_state.cv_a_valider)
    if nb:
        with st.container(key="carte_attention"):
            t, b = st.columns([5, 1.5], vertical_alignment="center")
            with t:
                ui.afficher(f'<div style="color:#3A4357;font-size:0.98rem"><b style="color:#0F1B3D">'
                            f'{nb} CV {"demande" if nb == 1 else "demandent"} votre attention.</b> '
                            f'Confirmez ou corrigez {"son" if nb == 1 else "leur"} domaine pour que CVSmart '
                            f'apprenne de votre choix.</div>')
            with b:
                st.button("Ouvrir la validation", key="btn_ouvrir_validation", type="primary",
                          use_container_width=True, on_click=aller_a, args=("Validation",))

# ============================================================
# PAGE FILTRER
# ============================================================
NIVEAUX = ['Bac', 'BTS/DUT', 'Licence', 'Master/Ingénieur', 'Doctorat']


def reinitialiser_filtres():
    for cle in list(st.session_state.keys()):
        if cle.startswith('f_'):
            del st.session_state[cle]


def enregistrer_selection(selection, etiquettes):
    nom = "_".join(etiquettes or ["Selection"])
    nom = re.sub(r'[^\w\-]+', '_', nom)[:80] + f"_{datetime.now().strftime('%Y%m%d_%H%M%S')}"
    dossier = os.path.join(DOSSIER_FILTRES, nom)
    os.makedirs(dossier, exist_ok=True)
    copies = 0
    for r in selection:
        for source in [r.get('chemin'), r.get('copie')]:
            if source and os.path.exists(source):
                shutil.copy2(source, os.path.join(dossier, r['fichier']))
                copies += 1
                break
    flash(f"{copies} CV copiés dans <b>{escape(dossier)}</b>")


def page_filtrer():
    resultats = [r for r in st.session_state.resultats if r.get('succes')]
    c1, c2 = st.columns([1.5, 1], vertical_alignment="center")
    with c1: ui.entete("Filtrer les CV")
    with c2:
        recherche = st.text_input("Recherche", placeholder="Nom, compétence, école...", key="f_recherche",
                                  label_visibility="collapsed")
    if not resultats:
        etat_vide("Aucun CV analysé. Lancez d'abord une analyse.", "Analyser des CV", "Analyser")
        return

    gauche, droite = st.columns([1, 2.3], gap="medium")
    with gauche, st.container(key="carte_filtres"):
        compteur = {}
        for r in resultats:
            compteur[r['categorie']] = compteur.get(r['categorie'], 0) + 1
        ui.afficher('<div class="cv-h3">Domaine</div>')
        domaines = [d for d, _ in sorted(compteur.items(), key=lambda x: -x[1])
                    if st.checkbox(d, key=f"f_dom_{nom_dossier(d)}")]

        exp_min = st.session_state.get('f_exp_min', 0)
        ui.afficher(f'<div class="cv-h3" style="margin-top:0.6rem">Expérience minimum : '
                    f'<span style="color:{ui.ROUGE}">{exp_min} an{"s" if exp_min > 1 else ""}</span></div>')
        exp_max = max([10] + [r['infos'].get('annees_experience', 0) for r in resultats])
        exp_min = st.slider("Expérience minimum", 0, exp_max, key="f_exp_min", label_visibility="collapsed")

        ui.afficher('<div class="cv-h3">Niveau d\'études</div>')
        niveau = st.selectbox("Niveau d'études", ["Tous les niveaux"] + NIVEAUX, key="f_niveau",
                              label_visibility="collapsed")

        ui.afficher('<div class="cv-h3">Compétences</div>')
        toutes_comp = sorted({c for r in resultats for c in r['infos'].get('competences', [])})
        competences = st.multiselect("Compétences", toutes_comp, key="f_comp", placeholder="+ Ajouter",
                                     label_visibility="collapsed")

        with st.expander("Plus de critères"):
            sexes = st.multiselect("Sexe", ['Homme', 'Femme', 'RAS'], key="f_sexe")
            villes = st.multiselect("Ville", sorted({r['infos'].get('ville', 'RAS') for r in resultats}), key="f_ville")
            ages = st.slider("Âge", 16, 70, (16, 70), key="f_age")

        st.button("Réinitialiser", key="btn_reinit_filtres", use_container_width=True, on_click=reinitialiser_filtres)

    selection = []
    for r in resultats:
        i = r['infos']
        if domaines and r['categorie'] not in domaines: continue
        if i.get('annees_experience', 0) < exp_min: continue
        if niveau != "Tous les niveaux" and i.get('niveau_etudes') != niveau: continue
        if not set(competences) <= set(i.get('competences', [])): continue
        if sexes and i.get('sexe') not in sexes: continue
        if villes and i.get('ville') not in villes: continue
        age = str(i.get('age', 'RAS'))
        if age.isdigit() and not ages[0] <= int(age) <= ages[1]: continue
        if recherche:
            cible = " ".join([i.get('nom', ''), i.get('prenom', ''), i.get('etablissement', ''),
                              " ".join(i.get('competences', [])), r['fichier'], i.get('texte_brut', '')]).lower()
            if recherche.lower() not in cible: continue
        selection.append(r)
    selection.sort(key=lambda r: r['confiance'], reverse=True)

    with droite:
        puces = [ui.pastille_domaine(d) for d in domaines]
        if exp_min: puces.append(ui.pastille(f"≥ {texte_experience(exp_min)} d'expérience", 'rose'))
        if niveau != "Tous les niveaux": puces.append(ui.pastille(niveau, 'bleu'))
        puces += [ui.pastille(c, 'gris') for c in competences]
        ui.afficher(f'<div style="display:flex;flex-wrap:wrap;gap:0.5rem;align-items:center;margin-bottom:0.2rem">'
                    f'<span style="color:#3A4357"><b style="color:{ui.BLEU}">{len(selection)} CV</b> sur {len(resultats)}</span>'
                    f'{"".join(puces)}</div>')
        if selection:
            lignes = ''.join(
                f'<tr><td><div class="cv-nom">{escape(nom_court(r["infos"]))}</div><div class="cv-fichier">{escape(r["fichier"])}</div></td>'
                f'<td>{ui.pastille_domaine(r["categorie"])}</td>'
                f'<td>{texte_experience(r["infos"].get("annees_experience", 0))}</td>'
                f'<td>{escape(r["infos"].get("niveau_etudes", "RAS")).replace("RAS", "—")}</td>'
                f'<td class="cv-droite">{r["confiance"] * 100:.0f} %</td></tr>' for r in selection)
            ui.afficher(f'<table class="cv-table"><thead><tr><th>Candidat</th><th>Domaine</th><th>Expérience</th>'
                        f'<th>Formation</th><th style="text-align:right">Confiance</th></tr></thead><tbody>{lignes}</tbody></table>')
        else:
            ui.message("Aucun CV ne correspond à ces critères.", 'ambre', 'alerte')

        _, b1, b2 = st.columns([1, 1.3, 1.4])
        with b1:
            st.button("Enregistrer ce filtre", key="btn_enregistrer_filtre", use_container_width=True,
                      disabled=not selection, on_click=enregistrer_selection, args=(selection, domaines),
                      help="Copie les CV de la sélection dans un dossier de CV_filtres")
        with b2:
            bouton_export("Exporter la sélection", selection, "btn_export_selection", "Selection")

# ============================================================
# PAGE VALIDATION
# ============================================================
CHOIX_VIDE = "Sélectionner..."
CHOIX_NOUVEAU = "Nouveau domaine..."


def selectionner_cv(idx):
    st.session_state.validation_sel = idx


def choisir_suggestion(idx, domaine):
    st.session_state[f"val_choix_{idx}"] = domaine
    st.session_state[f"val_autre_{idx}"] = CHOIX_VIDE


def cv_suivant(idx):
    file = st.session_state.cv_a_valider
    if idx in file and len(file) > 1:
        st.session_state.validation_sel = file[(file.index(idx) + 1) % len(file)]


def valider_domaine(idx):
    autre = st.session_state.get(f"val_autre_{idx}", CHOIX_VIDE)
    if autre == CHOIX_NOUVEAU:
        final = st.session_state.get(f"val_nouveau_{idx}", "").strip()
    elif autre != CHOIX_VIDE:
        final = autre
    else:
        final = st.session_state.get(f"val_choix_{idx}")
    if not final:
        flash("Choisissez un domaine avant de valider.", 'ambre', 'alerte')
        return
    r = st.session_state.resultats[idx]
    app = st.session_state.apprentissage
    apprendre = st.session_state.get(f"val_apprendre_{idx}", app['options'].get('auto', True))
    enregistrer_decision(app, r, final, apprendre)
    if final != r.get('categorie_predite') and idx not in st.session_state.cv_corriges:
        st.session_state.cv_corriges.append(idx)
    changer_domaine(r, final)
    cv_suivant(idx)
    st.session_state.cv_a_valider.remove(idx)
    if apprendre and app['options'].get('auto', True):
        reentrainer_classificateur()
    flash(f"<b>{escape(r['fichier'])}</b> validé en <b>{escape(final)}</b>.")


def domaines_disponibles():
    appris = {e['domaine'] for e in st.session_state.apprentissage.get('exemples', [])}
    return sorted((set(CATEGORIES) | appris) - {'Non classé'})


def page_validation():
    resultats = st.session_state.resultats
    file = [i for i in st.session_state.cv_a_valider if i < len(resultats)]
    nb_valides = sum(1 for i, r in enumerate(resultats) if r.get('succes') and i not in file)
    c1, c2 = st.columns([1.5, 1.2], vertical_alignment="center")
    with c1: ui.entete("Validation des classements")
    with c2:
        compteurs = [("À valider", len(file), 'ambre'), ("Validés", nb_valides, 'vert'),
                     ("Corrigés", len(st.session_state.cv_corriges), 'bleu')]
        ui.afficher('<div class="cv-compteurs">' + ''.join(
            f'<div class="cv-compteur" style="background:{ui.TONS[t][0]};color:{ui.TONS[t][1]}">{l} <b>{n}</b></div>'
            for l, n, t in compteurs) + '</div>')

    if not file:
        etat_vide("Aucun CV en attente de validation.", "Voir les résultats", "Résultats")
        return
    sel = st.session_state.validation_sel if st.session_state.validation_sel in file else file[0]
    r = resultats[sel]
    seuil = st.session_state.seuil_confiance

    gauche, droite = st.columns([1, 1.6], gap="medium")
    with gauche, st.container(key="carte_file"):
        ui.afficher('<div class="cv-h3">File d\'attente</div>')
        for idx in file:
            ri = resultats[idx]
            if ri['categorie'] == 'Non classé':
                lisible = len(ri['infos'].get('texte_brut', '').strip()) >= 300
                info = "Non classé · " + ("domaine incertain" if lisible else "texte peu lisible")
            else:
                info = f"{ri['categorie']} · confiance {ri['confiance'] * 100:.0f} %"
            with st.container(key=f"clic_file_{idx}"):
                ui.afficher(f'<div class="cv-file {"actif" if idx == sel else ""}"><div class="cv-file-nom">'
                            f'{escape(ri["fichier"])}</div><div class="cv-file-info">{escape(info)}</div></div>')
                st.button("Ouvrir", key=f"file_{idx}", on_click=selectionner_cv, args=(idx,))
        ui.afficher(f'<div class="cv-muted" style="margin-top:1rem">Les CV sous {seuil} % de confiance '
                    f'arrivent ici automatiquement.</div>')

    with droite, st.container(key="carte_detail"):
        cat = r['categorie']
        sous = ("Aucun domaine détecté avec certitude" if cat == 'Non classé'
                else f"Classé en {cat} avec {r['confiance'] * 100:.0f} % de confiance")
        h1, h2 = st.columns([2.5, 1.2])
        with h1:
            ui.afficher(f'<div style="display:flex;gap:1rem;align-items:center">'
                        f'<div class="cv-stat-icone" style="background:{ui.TONS["rose"][0]};width:56px;height:56px">'
                        f'{ui.icone("fichier", 24, ui.ROUGE)}</div><div><div style="font-size:1.35rem;font-weight:800;'
                        f'color:{ui.BLEU};word-break:break-all">{escape(r["fichier"])}</div><div class="cv-muted">{escape(sous)}</div></div></div>')
        with h2:
            if st.button("Modifier les infos", key="btn_modifier_infos", type="tertiary",
                         help="Corriger le nom, le contact, la formation… extraits du CV"):
                modale_correction(sel)
        ui.afficher('<div class="cv-h2" style="font-size:1.3rem;margin-top:0.2rem">Domaines suggérés</div>')

        suggestions = st.session_state.classif.suggerer(r['infos'].get('texte_brut', ''))
        if cat != 'Non classé' and cat not in [d for d, _ in suggestions]:
            suggestions = [(cat, r['confiance'])] + suggestions[:2]
        cle_choix = f"val_choix_{sel}"
        if cle_choix not in st.session_state:
            st.session_state[cle_choix] = cat if cat != 'Non classé' else (suggestions[0][0] if suggestions else None)
        for j, (dom, proba) in enumerate(suggestions):
            actif = st.session_state[cle_choix] == dom and st.session_state.get(f"val_autre_{sel}", CHOIX_VIDE) == CHOIX_VIDE
            couleur = ui.BLEU if j == 0 else "#5B6478"
            with st.container(key=f"clic_sugg_{sel}_{j}"):
                ui.afficher(f'<div class="cv-suggestion {"actif" if actif else ""}"><span class="cv-radio"></span>'
                            f'<span class="cv-suggestion-nom">{escape(dom)}</span>{ui.barre(proba * 100, couleur)}'
                            f'<span class="cv-suggestion-pct">{proba * 100:.0f} %</span></div>')
                st.button(f"Choisir {dom}", key=f"sugg_{sel}_{j}", on_click=choisir_suggestion, args=(sel, dom))

        ui.afficher('<div class="cv-h2" style="font-size:1.3rem;margin-top:0.4rem">Ou choisir un autre domaine</div>')
        options = [CHOIX_VIDE] + domaines_disponibles()
        if st.session_state.apprentissage['options'].get('nouveau_domaine'):
            options.append(CHOIX_NOUVEAU)
        s1, _ = st.columns([2.2, 1])
        with s1:
            autre = st.selectbox("Autre domaine", options, key=f"val_autre_{sel}", label_visibility="collapsed")
            if autre == CHOIX_NOUVEAU:
                st.text_input("Nom du nouveau domaine", key=f"val_nouveau_{sel}")

        a1, a2, _ = st.columns([1.6, 1.4, 0.4], vertical_alignment="center")
        with a1:
            st.checkbox("Apprendre de cette décision", key=f"val_apprendre_{sel}",
                        value=st.session_state.apprentissage['options'].get('auto', True))
        with a2:
            st.button("(voir l'apprentissage)", key="btn_voir_apprentissage", type="tertiary",
                      on_click=aller_a, args=("Apprentissage",))

        ui.afficher('<div style="height:2.5rem"></div>')
        _, b1, b2 = st.columns([0.5, 1.2, 1.2], vertical_alignment="center")
        with b1:
            st.button("Passer au suivant", key="btn_passer", use_container_width=True,
                      on_click=cv_suivant, args=(sel,), disabled=len(file) < 2)
        with b2:
            st.button("Valider le domaine", key="btn_valider_domaine", type="primary",
                      use_container_width=True, on_click=valider_domaine, args=(sel,))


@st.dialog("Correction des informations", width="large")
def modale_correction(idx_cv):
    r = st.session_state.resultats[idx_cv]
    infos = r['infos']
    st.markdown(f"**{r['fichier']}**")

    st.markdown("##### Identité")
    c1, c2 = st.columns(2)
    new_nom = c1.text_input("Nom", value=infos.get('nom', 'RAS'), key=f"cn_{idx_cv}")
    new_prenom = c2.text_input("Prénom", value=infos.get('prenom', 'RAS'), key=f"cp_{idx_cv}")
    c1, c2 = st.columns(2)
    new_email = c1.text_input("Email", value=infos.get('email', 'RAS'), key=f"ce_{idx_cv}")
    new_tel = c2.text_input("Téléphone", value=infos.get('telephone', 'RAS'), key=f"ct_{idx_cv}")
    options_sexe = ['RAS', 'Homme', 'Femme']
    sexe_actuel = infos.get('sexe', 'RAS')
    new_sexe = st.radio("Sexe", options_sexe, horizontal=True, key=f"cs_{idx_cv}",
                        index=options_sexe.index(sexe_actuel) if sexe_actuel in options_sexe else 0)
    c1, c2 = st.columns(2)
    age_actuel = str(infos.get('age', 'RAS'))
    # Âge inconnu → champ vide (évite d'enregistrer un âge inventé) ; borné pour éviter une exception Streamlit
    age_def = min(max(int(age_actuel), 15), 85) if age_actuel.isdigit() else None
    new_age = c1.number_input("Âge", min_value=15, max_value=85, value=age_def, key=f"ca_{idx_cv}")
    situations = ['RAS', 'Célibataire', 'Marié(e)', 'Divorcé(e)', 'Veuf(ve)']
    sit = infos.get('situation', 'RAS')
    new_situation = c2.selectbox("Situation", situations, key=f"csit_{idx_cv}",
                                 index=situations.index(sit) if sit in situations else 0)
    c1, c2 = st.columns(2)
    new_nationalite = c1.text_input("Nationalité", value=infos.get('nationalite', 'RAS'), key=f"cnat_{idx_cv}")
    new_ville = c2.text_input("Ville", value=infos.get('ville', 'RAS'), key=f"cv_{idx_cv}")

    st.markdown("##### Formation et expérience")
    niveaux = ['RAS'] + NIVEAUX
    niv = infos.get('niveau_etudes', 'RAS')
    c1, c2 = st.columns(2)
    new_niveau = c1.selectbox("Niveau d'études", niveaux, key=f"cniv_{idx_cv}",
                              index=niveaux.index(niv) if niv in niveaux else 0)
    new_etab = c2.text_input("Établissement", value=infos.get('etablissement', 'RAS'), key=f"cetab_{idx_cv}")
    c1, c2 = st.columns(2)
    new_experience = c1.number_input("Années d'expérience", min_value=0, max_value=60, step=1, key=f"cexp_{idx_cv}",
                                     value=min(max(int(infos.get('annees_experience', 0)), 0), 60))
    new_poste = c2.text_input("Poste actuel", value=infos.get('poste_actuel', 'RAS'), key=f"cposte_{idx_cv}")

    st.markdown("##### Compétences et langues")
    new_comps = st.text_area("Compétences (séparées par des virgules)", height=80, key=f"ccomp_{idx_cv}",
                             value=', '.join(infos.get('competences', [])))
    c1, c2 = st.columns(2)
    new_langues = c1.text_input("Langues", value=infos.get('langues', 'RAS'), key=f"clang_{idx_cv}")
    new_certs = c2.text_input("Certifications", value=infos.get('certifications', 'RAS'), key=f"ccert_{idx_cv}")

    st.markdown("##### Métier")
    metiers = ['RAS'] + list(CATEGORIES.get(r['categorie'], {}).get('metiers', {}).keys())
    if r.get('metier', 'RAS') not in metiers:
        metiers.insert(1, r['metier'])
    new_metier = st.selectbox(f"Métier ({r['categorie']})", metiers, key=f"cm_{idx_cv}",
                              index=metiers.index(r.get('metier', 'RAS')))

    c1, c2 = st.columns(2)
    if c1.button("Enregistrer", type="primary", use_container_width=True, key=f"save_{idx_cv}"):
        corrections = {
            'nom': new_nom, 'prenom': new_prenom, 'email': new_email, 'telephone': new_tel,
            'sexe': new_sexe, 'age': str(int(new_age)) if new_age is not None else 'RAS',
            'situation': new_situation, 'nationalite': new_nationalite, 'ville': new_ville,
            'niveau_etudes': new_niveau, 'etablissement': new_etab, 'annees_experience': new_experience,
            'poste_actuel': new_poste, 'langues': new_langues, 'certifications': new_certs,
        }
        app = st.session_state.apprentissage
        modifie = False
        for champ, val in corrections.items():
            if infos.get(champ, 'RAS') != val:
                enregistrer_correction(app, champ, infos.get(champ, 'RAS'), val, new_prenom, sauvegarder=False)
                modifie = True
            infos[champ] = val
        if modifie:
            sauvegarder_apprentissage(app)  # une seule écriture du fichier
        infos['competences'] = [c.strip() for c in new_comps.split(',') if c.strip()]
        r['metier'] = new_metier
        r['scores'] = calculer_score(infos)
        r['resume'] = generer_resume(infos, r['categorie'], new_metier)
        flash("Informations du candidat enregistrées.")
        st.rerun()
    if c2.button("Fermer", use_container_width=True, key=f"close_{idx_cv}"):
        st.rerun()

# ============================================================
# PAGE APPRENTISSAGE
# ============================================================
def relancer_apprentissage():
    reentrainer_classificateur()
    nb = len(st.session_state.apprentissage['exemples'])
    flash(f"Modèle réentraîné avec {nb} CV validé{'s' if nb > 1 else ''} par un humain.")


def sauver_options():
    app = st.session_state.apprentissage
    app['options']['auto'] = st.session_state.opt_auto
    app['options']['nouveau_domaine'] = st.session_state.opt_nouveau
    sauvegarder_apprentissage(app)


def page_apprentissage():
    app = st.session_state.apprentissage
    c1, c2 = st.columns([3, 1.3], vertical_alignment="center")
    with c1: ui.entete("Apprentissage du modèle")
    with c2:
        st.button("Relancer l'apprentissage", key="btn_relancer", type="primary", use_container_width=True,
                  on_click=relancer_apprentissage)

    s = stats_decisions(app)
    if s['precision'] is None:
        carte_precision = ui.carte_stat("Précision actuelle", "—", ui.BLEU, sous_texte="aucune décision pour le moment")
    else:
        couleur = ui.VERT if s['precision'] >= 80 else (ui.TONS['ambre'][1] if s['precision'] >= 60 else ui.ROUGE)
        carte_precision = ui.carte_stat("Précision actuelle", f"{s['precision']:.0f} %", couleur,
                                        sous_texte=f"sur {s['nb_decisions']} CV vérifiés", couleur_sous=couleur)
    maj = app['stats'].get('derniere_maj')
    heure_maj = datetime.fromisoformat(maj).strftime("%H:%M") if maj else "—"
    date_maj = datetime.fromisoformat(maj).strftime("%d/%m/%Y") if maj else ""
    ui.afficher(ui.grille([
        carte_precision,
        ui.carte_stat("Corrections apprises", s['corrections'], sous_texte=f"dont {s['corrections_jour']} aujourd'hui"),
        ui.carte_stat("CV d'entraînement", len(app['exemples']), sous_texte="validés par un humain"),
        ui.carte_stat("Dernière mise à jour", heure_maj,
                      sous_texte=f"Modèle v1.{app['stats'].get('version_modele', 0)} · {date_maj}".rstrip(' ·')),
    ], 4))

    gauche, droite = st.columns([1.55, 1], gap="medium")
    with gauche:
        if s['par_domaine']:
            pire = min(s['par_domaine'], key=lambda x: x[1])[0] if len(s['par_domaine']) > 1 else None
            lignes = ''.join(
                f'<div class="cv-precision"><span>{escape(d)}</span>{ui.barre(p, ui.ROUGE if d == pire else ui.BLEU)}'
                f'<span class="pct" style="color:{ui.ROUGE if d == pire else ui.BLEU}">{p:.0f} %</span></div>'
                for d, p, n in s['par_domaine'][:8])
        else:
            lignes = '<div class="cv-muted">Validez des CV pour mesurer la précision de chaque domaine.</div>'
        note = ''
        if s['confusion']:
            reel, predit = s['confusion']
            note = (f'<div class="cv-muted" style="margin-top:1.4rem">{escape(reel)} est le domaine le plus souvent '
                    f'confondu avec {escape(predit)} : vos corrections l\'aident à progresser.</div>')
        ui.afficher(f'<div class="cv-carte" style="min-height:420px"><div class="cv-h2" style="margin-bottom:0.8rem">'
                    f'Précision par domaine</div>{lignes}{note}</div>')

    with droite, st.container(key="carte_corrections"):
        if s['recentes']:
            liste = ''.join(
                f'<div class="cv-correction"><div style="font-weight:600;color:#0F1B3D">{escape(d["fichier"])}</div>'
                f'<div style="margin-top:0.2rem"><s>{escape(d["predit"])}</s> → <b style="color:{ui.BLEU}">'
                f'{escape(d["final"])}</b></div></div>' for d in s['recentes'])
        else:
            liste = '<div class="cv-muted">Aucune correction pour le moment.</div>'
        ui.afficher(f'<div class="cv-h2" style="margin-bottom:0.4rem">Corrections récentes</div>{liste}'
                    f'<div style="height:1.5rem"></div>')
        for libelle, cle, option in [("Apprentissage automatique", "opt_auto", 'auto'),
                                     ("Proposer un domaine nouveau", "opt_nouveau", 'nouveau_domaine')]:
            t, c = st.columns([5, 1], vertical_alignment="center")
            with t: ui.afficher(f'<b style="color:#0F1B3D">{libelle}</b>')
            with c:
                st.checkbox(libelle, value=app['options'].get(option, False), key=cle,
                            label_visibility="collapsed", on_change=sauver_options)

# ============================================================
# PAGE CORRESPONDANCE
# ============================================================
def page_correspondance():
    ui.entete("Correspondance avec une offre",
              "Collez une offre d'emploi pour classer les CV analysés selon leur adéquation.")
    with st.container(key="carte_offre"):
        offre = st.text_area("Description de l'offre", height=180, key="offre_text",
                             placeholder="Ex : Comptable confirmé, 5 ans d'expérience minimum, maîtrise de Sage et Excel")
        _, b = st.columns([3, 1.2])
        lancer = b.button("Lancer la correspondance", type="primary", key="btn_matching", use_container_width=True)

    resultats = [r for r in st.session_state.resultats if r.get('succes')]
    if lancer:
        if not offre.strip():
            ui.message("Collez d'abord une offre d'emploi.", 'ambre', 'alerte')
            return
        if not resultats:
            etat_vide("Aucun CV analysé. Lancez d'abord une analyse.", "Analyser des CV", "Analyser")
            return
        correspondances = [(r, calculer_matching(r['infos'], offre)) for r in resultats]
        st.session_state.correspondances = sorted(correspondances, key=lambda x: x[1]['score'], reverse=True)

    correspondances = st.session_state.get('correspondances')
    if not correspondances:
        return
    lignes = ''
    for r, m in correspondances[:20]:
        ton = 'vert' if m['score'] >= 70 else ('ambre' if m['score'] >= 40 else 'rose')
        forts = ' '.join(ui.pastille(c, 'gris') for c in m['communs'][:4]) or '—'
        manque = ', '.join(escape(c) for c in m['manquants'][:3]) or '—'
        lignes += (f'<tr><td><div class="cv-nom">{escape(nom_court(r["infos"]))}</div><div class="cv-fichier">{escape(r["fichier"])}</div></td>'
                   f'<td>{ui.pastille_domaine(r["categorie"])}</td><td>{texte_experience(r["infos"].get("annees_experience", 0))}</td>'
                   f'<td>{forts}</td><td class="cv-muted">{manque}</td>'
                   f'<td class="cv-droite"><span class="cv-pill" style="background:{ui.TONS[ton][0]};color:{ui.TONS[ton][1]}">'
                   f'{m["score"]:.0f} %</span></td></tr>')
    ui.afficher(f'<div class="cv-h2" style="margin:0.8rem 0 0.6rem 0">Meilleures correspondances</div>'
                f'<table class="cv-table"><thead><tr><th>Candidat</th><th>Domaine</th><th>Expérience</th>'
                f'<th>Points forts</th><th>À compléter</th><th style="text-align:right">Correspondance</th></tr></thead>'
                f'<tbody>{lignes}</tbody></table>')

# ============================================================
# PAGE EMAILS
# ============================================================
def traiter_email(em, gestionnaire):
    """Analyse les pièces jointes d'un email puis le déplace dans le dossier des emails traités"""
    nb = 0
    illisibles = []
    for pj in em['pieces_jointes']:
        r = analyser_cv(pj['chemin'], pj['nom_original'], source=em['email_expediteur'],
                        copie_nom=pj['nom_sauvegarde'], sous_dossier_metier=True)
        nb += r['succes']
        if not r['succes']:
            illisibles.append(pj['nom_original'])
    # Un email dont un CV n'a pas pu être lu (ex. ancien format .doc) reste dans la boîte
    if illisibles:
        return nb, False, ("CV illisible, email laissé dans la boîte de réception : " + ", ".join(illisibles)
                           + (" (format .doc non pris en charge : demandez un PDF ou un DOCX)"
                              if any(n.lower().endswith('.doc') for n in illisibles) else ""))
    ok, msg = gestionnaire.deplacer_email(em['id'])
    return nb, ok, msg


def _sans_accents(texte):
    return ''.join(c for c in unicodedata.normalize('NFD', texte.lower()) if unicodedata.category(c) != 'Mn')


def email_correspond(em, recherche):
    """Vrai si tous les mots recherchés apparaissent dans l'email (sans tenir compte des accents ni des majuscules)"""
    mots = _sans_accents(recherche).split()
    if not mots:
        return True
    cible = _sans_accents(" ".join([em.get('sujet', ''), em.get('expediteur', ''), em.get('date', '')]
                                   + [pj['nom_original'] for pj in em.get('pieces_jointes', [])]))
    return all(mot in cible for mot in mots)


def page_emails():
    ui.entete("Boîte mail RH", "Récupérez et analysez automatiquement les CV reçus par email.")
    creds = creds_messagerie()
    if not creds:
        ui.message("Identifiants de messagerie introuvables. Renseignez-les dans l'onglet <b>Paramètres</b> "
                   "(ou via <code>.streamlit/secrets.toml</code>).", 'rose', 'erreur')
        return
    serveur, port, email_address, mot_de_passe, dossier_traites = creds

    with st.container(key="carte_boite"):
        t, b1, b2 = st.columns([4, 1.6, 1], vertical_alignment="center")
        with t:
            ui.afficher(f'<div class="cv-h3" style="margin-bottom:0.2rem">Boîte connectée : {escape(email_address)}</div>'
                        f'<div class="cv-muted">Serveur {escape(str(serveur))}:{port} · emails traités déplacés dans '
                        f'« {escape(dossier_traites)} »</div>')
        verifier = b1.button("Vérifier les emails", type="primary", key="btn_verifier_emails", use_container_width=True)
        if b2.button("Vider", key="btn_vider_emails", use_container_width=True):
            st.session_state.pop('emails_recuperes', None)
            st.rerun()

    def gestionnaire():
        return GestionnaireEmail(serveur, port, email_address, mot_de_passe, dossier_traites)

    if verifier:
        with st.spinner("Connexion à la boîte mail..."):
            g = gestionnaire()
            emails, msg = g.recuperer_emails(max_emails=30)
            g.se_deconnecter()
        st.session_state.emails_recuperes = emails
        if not emails:
            ui.message(escape(msg), 'bleu', 'boite')

    emails = st.session_state.get('emails_recuperes')
    if emails is None:
        ui.message("Cliquez sur « Vérifier les emails » pour récupérer les CV reçus.", 'bleu', 'email')
        return
    if not emails:
        return

    nb_cv = sum(e['nb_cv'] for e in emails)
    ui.afficher(ui.grille([
        ui.carte_stat("Emails", len(emails), icone_nom='email', ton='bleu'),
        ui.carte_stat("CV reçus", nb_cv, icone_nom='trombone', ton='rose'),
        ui.carte_stat("À traiter", len(emails), icone_nom='boite', ton='ambre'),
    ], 3))

    recherche = st.text_input("Rechercher un email", key="recherche_emails",
                              placeholder="Rechercher par sujet, expéditeur, date ou nom de fichier...",
                              label_visibility="collapsed")
    tous_les_emails = emails
    emails = [em for em in emails if email_correspond(em, recherche)]
    if recherche.strip():
        ui.afficher(f'<div class="cv-muted"><b style="color:{ui.BLEU}">{len(emails)} email{"s" if len(emails) > 1 else ""}</b>'
                    f' sur {len(tous_les_emails)} pour « {escape(recherche.strip())} »</div>')
        if not emails:
            ui.message("Aucun email ne correspond à cette recherche.", 'ambre', 'alerte')
            return

    for em in emails:
        i = em['id']  # clé stable : l'ordre affiché change avec la recherche
        with st.container(key=f"carte_email_{i}"):
            t, b = st.columns([4, 1.4])
            with t:
                pjs = ''.join(f'<div class="cv-muted" style="display:flex;gap:0.4rem;align-items:center">'
                              f'{ui.icone("trombone", 15, ui.GRIS)}{escape(pj["nom_original"])} ({pj["taille"] / 1024:.0f} Ko)</div>'
                              for pj in em['pieces_jointes'])
                ui.afficher(f'<div class="cv-h3" style="margin:0">{escape(em["sujet"])}</div>'
                            f'<div class="cv-muted" style="margin:0.2rem 0 0.5rem 0">{escape(em["expediteur"])} · {escape(em["date"])}</div>{pjs}')
            with b:
                if st.button("Analyser les CV", type="primary", key=f"btn_analyser_{i}", use_container_width=True):
                    with st.spinner(f"Analyse de {em['nb_cv']} CV..."):
                        g = gestionnaire()
                        g.se_connecter()
                        nb, ok, msg = traiter_email(em, g)
                        g.se_deconnecter()
                    st.session_state.emails_recuperes = [e for e in tous_les_emails if e['id'] != em['id']]
                    if ok:
                        flash(f"{nb} CV analysé{'s' if nb > 1 else ''} · email déplacé dans « {escape(dossier_traites)} ».")
                    else:
                        flash(f"{nb} CV analysé{'s' if nb > 1 else ''}, mais l'email n'a pas pu être déplacé : {escape(msg)}", 'ambre', 'alerte')
                    st.rerun()
                if st.button("Ignorer", key=f"btn_ignorer_{i}", use_container_width=True):
                    st.session_state.emails_recuperes = [e for e in tous_les_emails if e['id'] != em['id']]
                    st.rerun()

    # Avec une recherche active, seuls les emails affichés sont analysés
    filtre = len(emails) < len(tous_les_emails)
    libelle = (f"Analyser les {len(emails)} emails affichés" if filtre and len(emails) > 1
               else "Analyser l'email affiché" if filtre else "Analyser tous les emails")
    if st.button(libelle, type="primary", key="btn_analyser_tous"):
        total, echecs = 0, 0
        with st.spinner("Analyse des CV reçus..."):
            g = gestionnaire()
            g.se_connecter()
            for em in emails:
                nb, ok, _ = traiter_email(em, g)
                total += nb
                echecs += not ok
            g.se_deconnecter()
        traites = {em['id'] for em in emails}
        st.session_state.emails_recuperes = [e for e in tous_les_emails if e['id'] not in traites]
        flash(f"{total} CV analysés et classés."
              + (f" {echecs} email(s) n'ont pas pu être déplacés." if echecs else ""), 'ambre' if echecs else 'vert')
        st.rerun()

# ============================================================
# PAGE PARAMÈTRES
# ============================================================
def vider_session():
    st.session_state.resultats = []
    st.session_state.cv_a_valider = []
    st.session_state.cv_corriges = []
    st.session_state.journal = []
    st.session_state.derniere_analyse = None
    st.session_state.pop('correspondances', None)
    flash("Session vidée : les résultats d'analyse ont été effacés.")


def reinitialiser_apprentissage():
    version = st.session_state.apprentissage['stats'].get('version_modele', 0)
    st.session_state.apprentissage = apprentissage_vide()
    # Le numéro de version continue (il ne repart pas à zéro)
    st.session_state.apprentissage['stats']['version_modele'] = version
    sauvegarder_apprentissage(st.session_state.apprentissage)
    reentrainer_classificateur()
    st.session_state.confirmer_reinit = False
    flash("Apprentissage réinitialisé.")


def page_parametres():
    ui.entete("Paramètres")
    with st.container(key="carte_seuil"):
        ui.afficher('<div class="cv-h2">Seuil de confiance</div><div class="cv-muted">En dessous de ce seuil, '
                    'un CV est envoyé dans la file de validation.</div>')
        st.session_state.seuil_confiance = st.slider("Seuil (%)", 30, 95, value=st.session_state.seuil_confiance)

    with st.container(key="carte_dossiers"):
        ui.afficher('<div class="cv-h2">Dossiers</div>')
        c1, c2, c3 = st.columns(3)
        c1.text_input("CV classés", value=os.path.abspath(DOSSIER_TRI), disabled=True, key="p_d1")
        c2.text_input("Exports Excel", value=os.path.abspath(DOSSIER_EXCEL), disabled=True, key="p_d2")
        c3.text_input("Sélections filtrées", value=os.path.abspath(DOSSIER_FILTRES), disabled=True, key="p_d3")

    with st.container(key="carte_messagerie"):
        ui.afficher('<div class="cv-h2">Messagerie</div>'
                    '<div class="cv-muted">Identifiants IMAP pour récupérer les CV reçus par email (page Emails). '
                    'Le mot de passe est stocké en clair dans un fichier local.</div>')
        cfg = st.session_state.config_email
        c1, c2 = st.columns(2)
        with c1:
            serveur = st.text_input("Serveur IMAP", value=cfg.get('imap_server', ''), placeholder="imap.example.com")
            email_address = st.text_input("Adresse email", value=cfg.get('email_address', ''))
            dossier_traites = st.text_input("Dossier des emails traités", value=cfg.get('dossier_traites', 'CV traités'))
        with c2:
            port = st.number_input("Port", value=int(cfg.get('imap_port') or 993), min_value=1, max_value=65535)
            mdp = st.text_input("Mot de passe d'application", value=cfg.get('app_password', ''), type="password")

        b1, b2 = st.columns([1.2, 1.2])
        if b1.button("Enregistrer", type="primary", key="btn_save_email", use_container_width=True):
            cfg.update({'imap_server': serveur, 'imap_port': int(port), 'email_address': email_address,
                        'app_password': mdp, 'dossier_traites': dossier_traites})
            sauvegarder_config_email(cfg)
            st.session_state.config_email = cfg
            flash("Identifiants de messagerie enregistrés.")
            st.rerun()
        if b2.button("Tester la connexion", key="btn_test_email", use_container_width=True):
            with st.spinner("Connexion au serveur IMAP..."):
                ok, msg = tester_connexion(serveur, int(port), email_address, mdp)
            if ok:
                ui.message("Connexion IMAP réussie.", 'vert', 'valide')
            else:
                ui.message(f"Connexion impossible : {escape(msg)}", 'rose', 'erreur')

    with st.container(key="carte_donnees"):
        ui.afficher('<div class="cv-h2">Données</div>')
        c1, c2 = st.columns(2)
        with c1:
            ui.afficher('<div class="cv-muted">Efface les résultats de la session en cours (les fichiers classés sont conservés).</div>')
            st.button("Vider la session", key="btn_vider_session", on_click=vider_session)
        with c2:
            ui.afficher('<div class="cv-muted">Supprime tout ce que le modèle a appris de vos corrections.</div>')
            confirmer = st.checkbox("Je confirme la réinitialisation", key="confirmer_reinit")
            st.button("Réinitialiser l'apprentissage", key="btn_reinit_app", disabled=not confirmer,
                      on_click=reinitialiser_apprentissage)

    nb_metiers = sum(len(d.get('metiers', {})) for d in CATEGORIES.values())
    ui.afficher(f'<div class="cv-h2" style="margin-top:0.6rem">Domaines configurés</div>'
                f'<div class="cv-muted" style="margin-bottom:0.4rem">{len(CATEGORIES) - 1} domaines et {nb_metiers} métiers</div>')
    for cat, data in CATEGORIES.items():
        if cat == 'Non classé': continue
        metiers = data.get('metiers', {})
        with st.expander(f"{cat} ({len(metiers)} métiers)"):
            for metier, mots in metiers.items():
                st.markdown(f"**{metier}** : {', '.join(mots[:5])}")

    ui.afficher(f'<div class="cv-carte" style="margin-top:1rem"><div class="cv-h2">À propos</div>'
                f'<div class="cv-muted">CVSmart v{VERSION_APP} · CIBLE RH Emploi — classement automatique de CV '
                f'(PDF, DOCX, scans via OCR), {len(CATEGORIES) - 1} domaines, validation humaine et apprentissage '
                f'à partir de vos corrections.</div></div>')

# ============================================================
# MAIN
# ============================================================
FONCTIONS_PAGES = {
    "Accueil": page_accueil, "Analyser": page_analyser, "Emails": page_emails,
    "Résultats": page_resultats, "Filtrer": page_filtrer, "Validation": page_validation,
    "Apprentissage": page_apprentissage, "Correspondance": page_correspondance,
    "Paramètres": page_parametres,
}


def main():
    ui.injecter_css()
    init_session()
    afficher_menu()
    if st.session_state.flash:
        texte, ton, icone_nom = st.session_state.flash
        st.session_state.flash = None
        ui.message(texte, ton, icone_nom)
    FONCTIONS_PAGES.get(st.session_state.page_active, page_accueil)()


if __name__ == "__main__":
    main()
